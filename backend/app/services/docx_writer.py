"""High-fidelity HTML → DOCX writer for the native document editor.

The editor (TipTap) stores documents as HTML: headings, paragraphs, nested lists, tables with
merged cells, images, inline marks (bold/italic/underline/strike/sup/sub, colour, highlight,
font family/size, links), Word comments and tracked revisions, header/footer regions. This module
maps each of those onto real WordprocessingML so an exported file reopens in Word/LibreOffice/
Google Docs looking like the editor — and re-imports into the editor without loss (the reader in
document_conversion.py understands everything written here).
"""

from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass, field, replace
from datetime import datetime, timezone
from typing import Any, Iterable

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Mm, Pt, RGBColor
from PIL import Image

PX_PER_INCH = 96
DEFAULT_FONT = "Arial"
DEFAULT_FONT_PT = 12
HEADING_PT = {1: 20, 2: 16, 3: 14, 4: 12, 5: 12, 6: 12}
# Typography shared with the editor sheet (.document-sheet in index.css) so PDF/DOCX match the page:
# body text leads at 28 px (21 pt), headings at 1.2 × their size, no implicit paragraph spacing.
BODY_LINE_PT = 21
HEADING_LINE_FACTOR = 1.2
ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
# Word's fixed highlight palette (w:highlight); any other background becomes run shading (w:shd).
HIGHLIGHT_BY_HEX = {
    "#ffff00": "yellow", "#00ff00": "green", "#00ffff": "cyan", "#ff00ff": "magenta",
    "#0000ff": "blue", "#ff0000": "red", "#000080": "darkBlue", "#008080": "darkCyan",
    "#008000": "darkGreen", "#800080": "darkMagenta", "#800000": "darkRed", "#808000": "darkYellow",
    "#808080": "darkGray", "#c0c0c0": "lightGray", "#000000": "black",
}
NAMED_COLORS = {
    "black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF",
    "yellow": "FFFF00", "orange": "FFA500", "purple": "800080", "gray": "808080", "grey": "808080",
}
BLOCK_TAGS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "table", "blockquote", "pre", "hr", "section", "div", "figure"}
BULLET_GLYPHS = ["•", "◦", "▪"]
NUMBER_FORMATS = ["decimal", "lowerLetter", "lowerRoman"]


# ── small parsers ───────────────────────────────────────────────────────────────────────────

def parse_style(raw: str | None) -> dict[str, str]:
    out: dict[str, str] = {}
    for declaration in (raw or "").split(";"):
        if ":" in declaration:
            name, value = declaration.split(":", 1)
            out[name.strip().lower()] = value.strip()
    return out


def parse_color(value: str | None) -> str | None:
    """CSS colour → 'RRGGBB' (None for transparent/unknown)."""
    if not value:
        return None
    value = value.strip().lower()
    if value in NAMED_COLORS:
        return NAMED_COLORS[value]
    match = re.fullmatch(r"#([0-9a-f]{3}|[0-9a-f]{6})", value)
    if match:
        digits = match.group(1)
        if len(digits) == 3:
            digits = "".join(ch * 2 for ch in digits)
        return digits.upper()
    match = re.fullmatch(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)(?:\s*,\s*([\d.]+))?\s*\)", value)
    if match:
        if match.group(4) is not None and float(match.group(4)) == 0:
            return None
        return "".join(f"{min(255, int(match.group(i))):02X}" for i in (1, 2, 3))
    return None


def parse_font_size_pt(value: str | None) -> float | None:
    if not value:
        return None
    match = re.fullmatch(r"\s*([\d.]+)\s*(pt|px|em|rem)?\s*", value.lower())
    if not match:
        return None
    number, unit = float(match.group(1)), match.group(2) or "px"
    if unit == "pt":
        return number
    if unit == "px":
        return number * 0.75
    return number * DEFAULT_FONT_PT


def parse_length_px(value: str | None) -> float | None:
    if not value:
        return None
    match = re.fullmatch(r"\s*([\d.]+)\s*(px|pt|in|cm|mm|%)?\s*", str(value).lower())
    if not match or match.group(2) == "%":
        return None
    number, unit = float(match.group(1)), match.group(2) or "px"
    return {"px": 1, "pt": 4 / 3, "in": 96, "cm": 96 / 2.54, "mm": 96 / 25.4}[unit] * number


def clean_font_family(value: str | None) -> str | None:
    if not value:
        return None
    first = value.split(",")[0].strip().strip("'\"")
    return first or None


# ── formatting state ────────────────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class RunFormat:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    superscript: bool = False
    subscript: bool = False
    code: bool = False
    color: str | None = None
    background: str | None = None
    font: str | None = None
    size_pt: float | None = None
    href: str | None = None
    comments: tuple[tuple[str, str, str, str], ...] = ()  # (id, author, date, text)
    revision: tuple[str, str, str] | None = None  # (kind, author, date)


def derive_format(base: RunFormat, tag: Tag) -> RunFormat:
    name = tag.name
    fmt = base
    if name in {"strong", "b"}:
        fmt = replace(fmt, bold=True)
    elif name in {"em", "i"}:
        fmt = replace(fmt, italic=True)
    elif name == "u":
        fmt = replace(fmt, underline=True)
    elif name in {"s", "strike", "del"}:
        fmt = replace(fmt, strike=True)
    elif name == "sup":
        fmt = replace(fmt, superscript=True, subscript=False)
    elif name == "sub":
        fmt = replace(fmt, subscript=True, superscript=False)
    elif name == "code":
        fmt = replace(fmt, code=True)
    elif name == "a" and tag.get("href"):
        fmt = replace(fmt, href=str(tag.get("href")))
    elif name == "mark":
        fmt = replace(fmt, background=parse_color(parse_style(tag.get("style")).get("background-color")) or "FFFF00")

    if tag.get("data-docx-comment") == "true":
        entry = (
            str(tag.get("data-comment-id") or ""),
            str(tag.get("data-comment-author") or ""),
            str(tag.get("data-comment-date") or ""),
            str(tag.get("data-comment-text") or ""),
        )
        if entry not in fmt.comments:
            fmt = replace(fmt, comments=fmt.comments + (entry,))
    if tag.get("data-revision-kind"):
        fmt = replace(fmt, revision=(
            str(tag.get("data-revision-kind")),
            str(tag.get("data-revision-author") or ""),
            str(tag.get("data-revision-date") or ""),
        ))

    style = parse_style(tag.get("style"))
    if name != "mark":
        color = parse_color(style.get("color"))
        if color:
            fmt = replace(fmt, color=color)
        background = parse_color(style.get("background-color"))
        if background and name not in {"td", "th", "table", "tr"}:
            fmt = replace(fmt, background=background)
    font = clean_font_family(style.get("font-family"))
    if font:
        fmt = replace(fmt, font=font)
    size = parse_font_size_pt(style.get("font-size"))
    if size:
        fmt = replace(fmt, size_pt=size)
    weight = style.get("font-weight", "")
    if weight in {"bold", "bolder"} or (weight.isdigit() and int(weight) >= 600):
        fmt = replace(fmt, bold=True)
    if style.get("font-style") == "italic":
        fmt = replace(fmt, italic=True)
    decoration = style.get("text-decoration", "") + " " + style.get("text-decoration-line", "")
    if "underline" in decoration:
        fmt = replace(fmt, underline=True)
    if "line-through" in decoration:
        fmt = replace(fmt, strike=True)
    return fmt


# ── writer ─────────────────────────────────────────────────────────────────────────────────

@dataclass
class _CommentAnchor:
    author: str
    date: str
    text: str
    runs: list[Any] = field(default_factory=list)


class DocxWriter:
    def __init__(self, *, title: str = "", margins_px: dict[str, Any] | None = None, page: dict[str, Any] | None = None) -> None:
        self.document = Document()
        self._revision_id = 1000
        self._comment_anchors: dict[str, _CommentAnchor] = {}
        self._bullet_abstract: int | None = None
        self._number_abstract: int | None = None
        self._last_link: tuple[Any, str] | None = None  # (w:hyperlink element, href) for run merging
        self._setup_page(margins_px, page)
        self._setup_styles()
        if title:
            self.document.core_properties.title = title

    # page + styles -------------------------------------------------------------------------
    def _setup_page(self, margins_px: dict[str, Any] | None, page: dict[str, Any] | None = None) -> None:
        section = self.document.sections[0]
        size = str((page or {}).get("size") or "a4").lower()
        landscape = str((page or {}).get("orientation") or "portrait").lower() == "landscape"
        width, height = (Inches(8.5), Inches(11)) if size == "letter" else (Mm(210), Mm(297))
        if landscape:
            width, height = height, width
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width = width
        section.page_height = height
        vertical = parse_length_px(str((margins_px or {}).get("vertical", ""))) or 56
        self._vertical_margin_px = vertical
        horizontal = parse_length_px(str((margins_px or {}).get("horizontal", ""))) or 56
        section.top_margin = section.bottom_margin = Inches(vertical / PX_PER_INCH)
        section.left_margin = section.right_margin = Inches(horizontal / PX_PER_INCH)

    def _setup_styles(self) -> None:
        styles = self.document.styles
        normal = styles["Normal"]
        normal.font.name = DEFAULT_FONT
        normal.font.size = Pt(DEFAULT_FONT_PT)
        self._set_east_asian_font(normal.element, DEFAULT_FONT)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        normal.paragraph_format.line_spacing = Pt(BODY_LINE_PT)
        # Same page-break rule as the editor: never leave a single line of a paragraph alone at a page edge.
        normal.paragraph_format.widow_control = True
        for level, size in HEADING_PT.items():
            heading = styles[f"Heading {level}"]
            heading.font.name = DEFAULT_FONT
            heading.font.size = Pt(size)
            heading.font.bold = True
            heading.font.italic = False
            heading.font.color.rgb = RGBColor(0x17, 0x15, 0x1B)
            self._set_east_asian_font(heading.element, DEFAULT_FONT)
            heading.paragraph_format.space_before = Pt(0)
            heading.paragraph_format.space_after = Pt(0)
            heading.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            heading.paragraph_format.line_spacing = Pt(round(size * HEADING_LINE_FACTOR, 1))
            heading.paragraph_format.keep_with_next = True  # the editor keeps a heading with its first lines

    @staticmethod
    def _set_east_asian_font(element: Any, font: str) -> None:
        r_pr = element.get_or_add_rPr()
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.append(fonts)
        for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attribute), font)
        # Theme font attributes (e.g. headings' majorHAnsi → Calibri Light) win over the explicit names.
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(attribute)) is not None:
                del fonts.attrib[qn(attribute)]

    @property
    def text_width_px(self) -> float:
        section = self.document.sections[0]
        return (section.page_width - section.left_margin - section.right_margin) / Emu(Inches(1)) * PX_PER_INCH

    # public ---------------------------------------------------------------------------------
    def write_header_card(self, header: dict[str, Any] | None) -> None:
        """Optional document hero (title / subtitle) set in the editor's header card."""
        if not isinstance(header, dict):
            return
        title = str(header.get("title") or "").strip()
        subtitle = str(header.get("subtitle") or "").strip()
        if title:
            paragraph = self.document.add_paragraph(style="Title")
            run = paragraph.add_run(title)
            run.font.name = DEFAULT_FONT
        if subtitle:
            paragraph = self.document.add_paragraph(style="Subtitle")
            paragraph.add_run(subtitle)

    def write_header_footer(self, config: dict[str, Any] | None, raw_html: str = "") -> None:
        """Repeating page header/footer (logo, text, "Pagina X di Y") set in the editor's page-band dialog.

        Imported Word header/footer regions already present in the HTML win over these bands.
        """
        if not isinstance(config, dict):
            return
        section = self.document.sections[0]
        for kind, target in (("header", section.header), ("footer", section.footer)):
            band = config.get(kind)
            if not isinstance(band, dict) or band.get("enabled") is not True:
                continue
            if f'data-docx-region="{kind}"' in raw_html:
                continue
            offset = parse_length_px(str(band.get("edgeOffset", ""))) if band.get("edgeOffset") is not None else None
            distance = Inches(min(1.2, max(0.0, (offset if offset is not None else 6) / PX_PER_INCH)))
            if kind == "header":
                section.header_distance = distance
            else:
                section.footer_distance = distance
            self._write_band(band, target, kind)

    @staticmethod
    def _band_items(band: dict[str, Any], kind: str) -> list[dict[str, Any]]:
        """Items of a band; bands saved before free positioning (logoUrl/text/showPageNumber) are converted."""
        raw = band.get("items")
        if isinstance(raw, list):
            return [item for item in raw if isinstance(item, dict)][:8]
        slots = {"left": 0.0, "center": 50.0, "right": 100.0}
        style = {"fontFamily": "Arial", "sizePt": 9, "color": "#64748b", "bold": False, "italic": False}
        items: list[dict[str, Any]] = []
        if band.get("logoUrl"):
            logo_x = band.get("logoX")
            x = float(logo_x) if isinstance(logo_x, (int, float)) and not isinstance(logo_x, bool) else slots.get(str(band.get("logoAlign")), 0.0)
            items.append({"type": "logo", "url": band.get("logoUrl"), "height": band.get("logoHeight") or 36, "x": x})
        if str(band.get("text") or "").strip():
            items.append({"type": "text", "text": band.get("text"), "x": slots.get(str(band.get("textAlign")), 100.0 if kind == "header" else 0.0), **style})
        if band.get("showPageNumber") is True:
            items.append({"type": "page", "x": slots.get(str(band.get("pageNumberAlign")), 100.0), **style})
        return items

    @staticmethod
    def _text_width_px(text: str, family: str, size_pt: float, bold: bool, italic: bool) -> float:
        """Width of a text run, measured with the font the exporter will actually use (fontconfig match)."""
        import subprocess

        from PIL import ImageFont

        size_px = size_pt * PX_PER_INCH / 72
        pattern = f"{family}:{'bold' if bold else 'regular'}{':italic' if italic else ''}"
        try:
            path = subprocess.run(["fc-match", "-f", "%{file}", pattern], capture_output=True, text=True, timeout=5, check=False).stdout.strip()
            if path:
                return float(ImageFont.truetype(path, max(1, round(size_px * 4))).getlength(text)) / 4
        except Exception:
            pass
        return len(text) * size_px * 0.52

    def _band_metrics(self, items: list[dict[str, Any]], band: dict[str, Any]) -> float:
        """Tallest allowed logo (px). Mirrors bandMetrics() in the editor (documentHeaderFooter.ts)."""
        sizes = []
        for raw in items:
            if raw.get("type") == "page" or (raw.get("type") == "text" and str(raw.get("text") or "").strip()):
                sizes.append(max(6.0, min(48.0, float(raw.get("sizePt") or 9))) * PX_PER_INCH / 72)
        descent = max(sizes) * 0.212 if sizes else 0.0
        try:
            offset = max(0.0, min(60.0, float(band.get("edgeOffset", 6))))
        except (TypeError, ValueError):
            offset = 6.0
        available = max(24.0, float(getattr(self, "_vertical_margin_px", 56)) - offset - 4 - 2)
        return max(12.0, float(int(available - descent)))

    def _write_band(self, band: dict[str, Any], target: Any, kind: str) -> None:
        width_px = self.text_width_px
        items_raw = self._band_items(band, kind)
        logo_max = self._band_metrics(items_raw, band)
        # (alignment, position px) -> items sharing that tab stop, in reading order.
        anchors: dict[tuple[float, str], list[tuple[str, Any]]] = {}

        for raw in items_raw:
            try:
                x = max(0.0, min(100.0, float(raw.get("x", 0))))
            except (TypeError, ValueError):
                x = 0.0
            kind_of = raw.get("type")
            style = {
                "family": clean_font_family(str(raw.get("fontFamily") or "")) or DEFAULT_FONT,
                "size": max(6.0, min(48.0, float(raw.get("sizePt") or 9))),
                "color": str(raw.get("color") or "#64748b"),
                "bold": raw.get("bold") is True,
                "italic": raw.get("italic") is True,
            }
            item_width = 0.0
            if kind_of == "logo":
                data = self._decode_data_image(raw.get("url"))
                if data is None:
                    continue
                try:
                    height_px = max(16.0, min(logo_max, float(raw.get("height") or 36)))
                except (TypeError, ValueError):
                    height_px = min(36.0, logo_max)
                with Image.open(io.BytesIO(data)) as image:
                    item_width = height_px * image.width / max(1, image.height)
                payload: Any = (data, height_px)
            elif kind_of == "text":
                text = str(raw.get("text") or "").strip()[:200]
                if not text:
                    continue
                item_width = self._text_width_px(text, style["family"], style["size"], style["bold"], style["italic"])
                payload = (text, style)
            elif kind_of == "page":
                item_width = self._text_width_px("Pagina 10 di 10", style["family"], style["size"], style["bold"], style["italic"])
                payload = (None, style)
            else:
                continue
            # Flush-right / centred items use right / centre tab stops, so their position does not depend on a
            # measured width (a flush-right "Pagina 15 di 16" must never wrap onto a second line).
            if x >= 99.5:
                key = (round(width_px, 1), "right")
            elif abs(x - 50) < 0.5:
                key = (round(width_px / 2, 1), "center")
            else:
                key = (round(max(0.0, width_px - item_width - 1) * x / 100, 1), "left")
            anchors.setdefault(key, []).append((str(kind_of), payload))
        if not anchors:
            return

        target.is_linked_to_previous = False
        paragraph = target.paragraphs[0] if target.paragraphs else target.add_paragraph()
        # "Header"/"Footer" styles ship their own centre/right tab stops, which would capture our tabs.
        paragraph.style = self.document.styles["Normal"]
        fmt = paragraph.paragraph_format
        fmt.space_after = Pt(0)
        fmt.space_before = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE  # same line box as the editor (font size × 1.15)
        self._shrink_paragraph_mark(paragraph)
        alignments = {"left": WD_TAB_ALIGNMENT.LEFT, "center": WD_TAB_ALIGNMENT.CENTER, "right": WD_TAB_ALIGNMENT.RIGHT}
        for (position, align), items in sorted(anchors.items(), key=lambda entry: (entry[0][0], entry[0][1])):
            if position > 0 or align != "left":
                fmt.tab_stops.add_tab_stop(Emu(int(Inches(position / PX_PER_INCH))), alignments[align])
                paragraph.add_run("\t")
            for index, (item, payload) in enumerate(items):
                if index:
                    paragraph.add_run("  ")
                if item == "logo":
                    data, height_px = payload
                    self._add_logo(paragraph, data, height_px)
                elif item == "text":
                    text, style = payload
                    self._styled_run(paragraph.add_run(text), style)
                else:
                    self._page_number_runs(paragraph, payload[1])
        if band.get("showRule") is not False:
            self._edge_border(paragraph, "bottom" if kind == "header" else "top")

    @staticmethod
    def _shrink_paragraph_mark(paragraph: Any) -> None:
        """The paragraph mark would add the Normal font's descent under a logo-only line: make it 1 pt."""
        p_pr = paragraph._p.get_or_add_pPr()
        r_pr = p_pr.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = OxmlElement("w:rPr")
            p_pr.append(r_pr)
        for tag in ("w:sz", "w:szCs"):
            element = OxmlElement(tag)
            element.set(qn("w:val"), "2")
            r_pr.append(element)

    @staticmethod
    def _add_logo(paragraph: Any, data: bytes, height_px: float) -> None:
        run = paragraph.add_run()
        run.add_picture(io.BytesIO(data), height=Inches(height_px / PX_PER_INCH))
        # python-docx omits the wrap distances; LibreOffice then assumes 0.125" (9 pt) around the image,
        # which shifts it right and makes a flush-right logo wrap onto a second line.
        for inline in run._r.iter(qn("wp:inline")):
            for attribute in ("distT", "distB", "distL", "distR"):
                inline.set(attribute, "0")

    @staticmethod
    def _styled_run(run: Any, style: dict[str, Any]) -> None:
        run.font.name = style["family"]
        run.font.size = Pt(style["size"])
        run.font.bold = style["bold"] or None
        run.font.italic = style["italic"] or None
        color = parse_color(style["color"])
        if color is not None:
            run.font.color.rgb = RGBColor.from_string(color)
        r_pr = run._r.get_or_add_rPr()
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.append(fonts)
        for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attribute), style["family"])

    @staticmethod
    def _decode_data_image(value: Any) -> bytes | None:
        match = re.match(r"data:image/(?:png|jpe?g|gif|webp);base64,(.+)", str(value or ""), re.DOTALL)
        if not match:
            return None
        try:
            data = base64.b64decode(match.group(1), validate=False)
            with Image.open(io.BytesIO(data)) as image:
                if image.format not in {"PNG", "JPEG", "GIF"}:
                    buffer = io.BytesIO()
                    image.convert("RGBA").save(buffer, format="PNG")
                    return buffer.getvalue()
            return data
        except Exception:
            return None

    def _field_run(self, paragraph: Any, instruction: str, placeholder: str, style: dict[str, Any]) -> None:
        run = paragraph.add_run()
        self._styled_run(run, style)
        for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, placeholder), ("end", None)):
            if kind:
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), kind)
            elif text == instruction:
                element = OxmlElement("w:instrText")
                element.set(qn("xml:space"), "preserve")
                element.text = f" {instruction} "
            else:
                element = OxmlElement("w:t")
                element.text = placeholder
            run._r.append(element)

    def _page_number_runs(self, paragraph: Any, style: dict[str, Any]) -> None:
        self._styled_run(paragraph.add_run("Pagina "), style)
        self._field_run(paragraph, "PAGE", "1", style)
        self._styled_run(paragraph.add_run(" di "), style)
        self._field_run(paragraph, "NUMPAGES", "1", style)

    @staticmethod
    def _edge_border(paragraph: Any, edge: str) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        side = OxmlElement(f"w:{edge}")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "3"), ("w:color", "CBD5E1")):
            side.set(qn(key), value)
        borders.append(side)
        # CT_PPr is an ordered sequence: pBdr must precede shd/tabs/spacing/ind/jc or Word may reject the part.
        successors = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
                      "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                      "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                      "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
        for child in p_pr:
            if child.tag in {qn(tag) for tag in successors}:
                child.addprevious(borders)
                return
        p_pr.append(borders)

    def write_html(self, raw_html: str, page_breaks: Iterable[int] | None = None) -> None:
        """Write the document body. ``page_breaks`` are indexes of top-level blocks that start a new page in the
        editor: forcing them keeps the exported pages identical to the sheets the user saw (LibreOffice and Word
        decide widow/orphan breaks differently from the editor)."""
        soup = BeautifulSoup(raw_html or "<p></p>", "html.parser")
        body = self.document
        breaks = {int(index) for index in (page_breaks or []) if isinstance(index, (int, float)) and index > 0}
        block_index = -1
        for node in soup.children:
            if isinstance(node, Tag):
                block_index += 1
            if isinstance(node, Tag) and node.name == "section" and node.get("data-docx-region"):
                self._write_region(node)
                continue
            before = len(body.paragraphs)
            self._write_node(node, body)
            if block_index in breaks and len(body.paragraphs) > before:
                body.paragraphs[before].paragraph_format.page_break_before = True
        self._flush_comments()

    def save(self) -> bytes:
        output = io.BytesIO()
        self.document.save(output)
        return output.getvalue()

    # regions ------------------------------------------------------------------------------
    def _write_region(self, node: Tag) -> None:
        kind = str(node.get("data-docx-region") or "header")
        section = self.document.sections[0]
        if kind.endswith("-first"):
            section.different_first_page_header_footer = True
        if kind.endswith("-even"):
            self.document.settings.odd_and_even_pages_header_footer = True
        target = {
            "header": section.header,
            "header-first": section.first_page_header,
            "header-even": section.even_page_header,
            "footer": section.footer,
            "footer-first": section.first_page_footer,
            "footer-even": section.even_page_footer,
        }.get(kind, section.header)
        target.is_linked_to_previous = False
        placeholder = target.paragraphs[0] if target.paragraphs and not target.paragraphs[0].text else None
        for child in node.children:
            self._write_node(child, target)
        if placeholder is not None and len(target.paragraphs) > 1:
            placeholder._element.getparent().remove(placeholder._element)

    # blocks -------------------------------------------------------------------------------
    def _write_node(self, node: Any, container: Any, base: RunFormat = RunFormat()) -> None:
        if isinstance(node, NavigableString):
            if str(node).strip():
                paragraph = container.add_paragraph()
                self._write_inline(node, paragraph, base)
            return
        if not isinstance(node, Tag):
            return
        name = node.name
        if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            paragraph = container.add_paragraph(style=f"Heading {name[1]}")
            self._apply_paragraph_style(paragraph, node)
            self._write_children_inline(node, paragraph, base)
        elif name == "p":
            paragraph = container.add_paragraph()
            self._apply_paragraph_style(paragraph, node)
            self._write_children_inline(node, paragraph, base)
        elif name in {"ul", "ol"}:
            self._write_list(node, container, level=0, base=base)
        elif name == "table":
            self._write_table(node, container, base)
        elif name == "blockquote":
            for child in self._block_children(node):
                if isinstance(child, Tag) and child.name == "p":
                    paragraph = container.add_paragraph(style=self._style_or_default("Quote"))
                    self._apply_paragraph_style(paragraph, child)
                    self._write_children_inline(child, paragraph, replace(base, italic=True))
                else:
                    self._write_node(child, container, base)
        elif name == "pre":
            for line in node.get_text().split("\n"):
                paragraph = container.add_paragraph(style=self._style_or_default("No Spacing"))
                run = paragraph.add_run(line)
                run.font.name = "Courier New"
                run.font.size = Pt(10)
        elif name == "hr":
            paragraph = container.add_paragraph()
            if "page-break" in " ".join(node.get("class") or []):
                paragraph.add_run().add_break(WD_BREAK.PAGE)
            else:
                self._bottom_border(paragraph)
        elif name == "img":
            paragraph = container.add_paragraph()
            self._write_image(node, paragraph)
        elif name == "br":
            container.add_paragraph()
        elif name == "span" and node.get("data-type") == "blockMath":
            paragraph = container.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._write_math(node, paragraph)
        elif name in {"section", "div", "figure", "tbody", "thead", "article", "main", "body", "html"}:
            has_blocks = any(isinstance(child, Tag) and child.name in BLOCK_TAGS for child in node.children)
            if has_blocks:
                for child in node.children:
                    self._write_node(child, container, derive_format(base, node))
            elif node.get_text(strip=True) or node.find("img"):
                paragraph = container.add_paragraph()
                self._apply_paragraph_style(paragraph, node)
                self._write_children_inline(node, paragraph, derive_format(base, node))
        else:
            # Stray inline content at block level: wrap in a paragraph.
            if node.get_text(strip=True) or node.find("img") or node.get("data-type"):
                paragraph = container.add_paragraph()
                self._write_inline(node, paragraph, base)

    @staticmethod
    def _block_children(node: Tag) -> Iterable[Any]:
        return [child for child in node.children if not (isinstance(child, NavigableString) and not str(child).strip())]

    def _style_or_default(self, name: str) -> str | None:
        try:
            self.document.styles[name]
            return name
        except KeyError:
            return None

    @staticmethod
    def _paragraph_font_pt(node: Tag) -> float:
        """Largest font size used by the paragraph (explicit run sizes, else its block default)."""
        sizes = [
            size
            for element in node.find_all(True)
            if (size := parse_font_size_pt(parse_style(element.get("style")).get("font-size"))) is not None
        ]
        own = parse_font_size_pt(parse_style(node.get("style")).get("font-size"))
        if own:
            sizes.append(own)
        if sizes:
            return max(sizes)
        if node.name and re.fullmatch(r"h[1-6]", node.name):
            return float(HEADING_PT[int(node.name[1])])
        return float(DEFAULT_FONT_PT)

    def _apply_paragraph_style(self, paragraph: Any, node: Tag) -> None:
        style = parse_style(node.get("style"))
        paragraph.paragraph_format.widow_control = True  # explicit: LibreOffice ignores the style-level flag
        alignment = ALIGNMENTS.get(style.get("text-align", "").lower())
        if alignment is not None:
            paragraph.alignment = alignment
        line_height = style.get("line-height")
        if line_height:
            base_pt = self._paragraph_font_pt(node)
            try:
                # CSS unitless line-height scales with the font size (not with the font's own metrics like
                # Word's "multiple"), so convert to an at-least value in points.
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                paragraph.paragraph_format.line_spacing = Pt(round(float(line_height) * base_pt, 1))
            except ValueError:
                size = parse_font_size_pt(line_height)
                if size:
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                    paragraph.paragraph_format.line_spacing = Pt(size)
        indent = parse_length_px(style.get("margin-left") or style.get("padding-left"))
        if indent:
            paragraph.paragraph_format.left_indent = Inches(indent / PX_PER_INCH)
        first_line = parse_length_px(style.get("text-indent"))
        if first_line:
            paragraph.paragraph_format.first_line_indent = Inches(first_line / PX_PER_INCH)
        space_before = parse_length_px(style.get("margin-top"))
        if space_before is not None:
            paragraph.paragraph_format.space_before = Pt(space_before * 0.75)
        space_after = parse_length_px(style.get("margin-bottom"))
        if space_after is not None:
            paragraph.paragraph_format.space_after = Pt(space_after * 0.75)

    @staticmethod
    def _bottom_border(paragraph: Any) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "BFBFBF")):
            bottom.set(qn(key), value)
        borders.append(bottom)
        p_pr.append(borders)

    # lists --------------------------------------------------------------------------------
    def _abstract_num(self, ordered: bool) -> int:
        cached = self._number_abstract if ordered else self._bullet_abstract
        if cached is not None:
            return cached
        numbering = self.document.part.numbering_part.element
        existing = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
        abstract_id = max(existing, default=0) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "hybridMultilevel")
        abstract.append(multi)
        for level in range(9):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            fmt = OxmlElement("w:numFmt")
            text = OxmlElement("w:lvlText")
            if ordered:
                fmt.set(qn("w:val"), NUMBER_FORMATS[level % 3])
                text.set(qn("w:val"), f"%{level + 1}.")
            else:
                fmt.set(qn("w:val"), "bullet")
                text.set(qn("w:val"), BULLET_GLYPHS[level % 3])
            jc = OxmlElement("w:lvlJc")
            jc.set(qn("w:val"), "left")
            p_pr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(720 * (level + 1)))
            ind.set(qn("w:hanging"), "360")
            p_pr.append(ind)
            for element in (start, fmt, text, jc, p_pr):
                lvl.append(element)
            if not ordered:
                r_pr = OxmlElement("w:rPr")
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), "Arial")
                fonts.set(qn("w:hAnsi"), "Arial")
                r_pr.append(fonts)
                lvl.append(r_pr)
            abstract.append(lvl)
        # abstractNum elements must precede num elements
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        if ordered:
            self._number_abstract = abstract_id
        else:
            self._bullet_abstract = abstract_id
        return abstract_id

    def _new_num(self, ordered: bool, start: int = 1) -> int:
        abstract_id = self._abstract_num(ordered)
        numbering = self.document.part.numbering_part.element
        existing = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
        num_id = max(existing, default=0) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        if ordered:
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), "0")
            start_override = OxmlElement("w:startOverride")
            start_override.set(qn("w:val"), str(max(1, start)))
            override.append(start_override)
            num.append(override)
        numbering.append(num)
        return num_id

    @staticmethod
    def _set_numbering(paragraph: Any, num_id: int, level: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(min(8, level)))
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num)
        p_pr.append(num_pr)

    def _write_list(self, node: Tag, container: Any, level: int, base: RunFormat) -> None:
        ordered = node.name == "ol"
        try:
            start = int(node.get("start") or 1)
        except ValueError:
            start = 1
        num_id = self._new_num(ordered, start)
        list_style = self._style_or_default("List Paragraph")
        for item in node.find_all("li", recursive=False):
            numbered = False
            inline_buffer: list[Any] = []

            def flush_inline() -> None:
                nonlocal numbered
                if not inline_buffer:
                    return
                paragraph = container.add_paragraph(style=list_style)
                if not numbered:
                    self._set_numbering(paragraph, num_id, level)
                    numbered = True
                else:
                    paragraph.paragraph_format.left_indent = Inches(0.5 * (level + 1))
                for piece in inline_buffer:
                    self._write_inline(piece, paragraph, base)
                inline_buffer.clear()

            for child in item.children:
                if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                    flush_inline()
                    self._write_list(child, container, level + 1, base)
                elif isinstance(child, Tag) and child.name in {"p", "div"}:
                    flush_inline()
                    paragraph = container.add_paragraph(style=list_style)
                    self._apply_paragraph_style(paragraph, child)
                    if not numbered:
                        self._set_numbering(paragraph, num_id, level)
                        numbered = True
                    else:
                        paragraph.paragraph_format.left_indent = Inches(0.5 * (level + 1))
                    self._write_children_inline(child, paragraph, base)
                elif isinstance(child, NavigableString) and not str(child).strip() and not inline_buffer:
                    continue  # layout whitespace between blocks inside <li>
                else:
                    inline_buffer.append(child)
            flush_inline()
            if not numbered:
                paragraph = container.add_paragraph(style=list_style)
                self._set_numbering(paragraph, num_id, level)

    # tables -------------------------------------------------------------------------------
    def _write_table(self, node: Tag, container: Any, base: RunFormat) -> None:
        rows = [row for row in node.find_all("tr") if row.find_parent("table") is node]
        if not rows:
            return
        grid: list[list[Tag | None]] = []
        spans: list[tuple[int, int, int, int, Tag]] = []  # r, c, rowspan, colspan, cell
        occupied: set[tuple[int, int]] = set()
        for r_index, row in enumerate(rows):
            c_index = 0
            for cell in row.find_all(["td", "th"], recursive=False):
                while (r_index, c_index) in occupied:
                    c_index += 1
                try:
                    colspan = max(1, int(cell.get("colspan") or 1))
                    rowspan = max(1, int(cell.get("rowspan") or 1))
                except ValueError:
                    colspan, rowspan = 1, 1
                spans.append((r_index, c_index, rowspan, colspan, cell))
                for dr in range(rowspan):
                    for dc in range(colspan):
                        occupied.add((r_index + dr, c_index + dc))
                c_index += colspan
        row_count = max(len(rows), max((r + rs for r, _, rs, _, _ in spans), default=1))
        col_count = max((c + cs for _, c, _, cs, _ in spans), default=1)
        try:
            table = container.add_table(rows=row_count, cols=col_count)
        except TypeError:  # header/footer containers also need a width
            table = container.add_table(rows=row_count, cols=col_count, width=Inches(self.text_width_px / PX_PER_INCH))
        try:
            table.style = self.document.styles["Table Grid"]
        except KeyError:
            pass
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        header_rows = {r for r, _, _, _, cell in spans if cell.name == "th"}
        for r_index in header_rows:
            if r_index == 0 or (r_index - 1) in header_rows:
                tr_pr = table.rows[r_index]._tr.get_or_add_trPr()
                repeat = OxmlElement("w:tblHeader")
                tr_pr.append(repeat)

        for r_index, c_index, rowspan, colspan, cell_node in spans:
            cell = table.cell(r_index, c_index)
            if rowspan > 1 or colspan > 1:
                cell = cell.merge(table.cell(min(row_count, r_index + rowspan) - 1, min(col_count, c_index + colspan) - 1))
            background = parse_color(parse_style(cell_node.get("style")).get("background-color"))
            if background:
                tc_pr = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), background)
                tc_pr.append(shading)
            placeholder = cell.paragraphs[0]
            cell_format = replace(base, bold=True) if cell_node.name == "th" else base
            has_blocks = any(isinstance(child, Tag) and child.name in BLOCK_TAGS for child in cell_node.children)
            if has_blocks:
                for child in cell_node.children:
                    self._write_node(child, cell, cell_format)
            elif cell_node.get_text(strip=True) or cell_node.find("img"):
                paragraph = cell.add_paragraph()
                self._write_children_inline(cell_node, paragraph, cell_format)
            if len(cell.paragraphs) > 1:
                placeholder._element.getparent().remove(placeholder._element)

    # inline -------------------------------------------------------------------------------
    def _write_children_inline(self, node: Tag, paragraph: Any, base: RunFormat) -> None:
        for child in node.children:
            self._write_inline(child, paragraph, base)

    def _write_inline(self, node: Any, paragraph: Any, fmt: RunFormat) -> None:
        if isinstance(node, NavigableString):
            text = str(node)
            if not text:
                return
            # Collapse HTML whitespace like a browser would (TipTap text has no layout newlines).
            text = re.sub(r"[ \t\r\n]+", " ", text)
            if text == " " and not paragraph.runs:
                return
            self._add_text_run(paragraph, text, fmt)
            return
        if not isinstance(node, Tag):
            return
        if node.name == "br":
            run = self._new_run(paragraph, fmt)
            run.add_break()
            return
        if node.name == "img":
            self._write_image(node, paragraph)
            return
        if node.name == "span" and node.get("data-type") in {"inlineMath", "blockMath"}:
            self._write_math(node, paragraph)
            return
        if node.name in {"ul", "ol", "table"}:
            # Block inside inline context (rare): append flattened text.
            self._add_text_run(paragraph, node.get_text(" ", strip=True), fmt)
            return
        child_format = derive_format(fmt, node)
        for child in node.children:
            self._write_inline(child, paragraph, child_format)

    def _write_math(self, node: Tag, paragraph: Any) -> None:
        latex = str(node.get("data-latex") or node.get_text() or "").strip()
        if not latex:
            return
        run = paragraph.add_run(latex)
        run.font.name = "Cambria Math"
        run.italic = True

    def _add_text_run(self, paragraph: Any, text: str, fmt: RunFormat) -> Any:
        run = self._new_run(paragraph, fmt)
        if fmt.revision and fmt.revision[0] == "delete":
            deleted = OxmlElement("w:delText")
            deleted.set(qn("xml:space"), "preserve")
            deleted.text = text
            run._r.append(deleted)
        else:
            run.text = text
        return run

    def _new_run(self, paragraph: Any, fmt: RunFormat) -> Any:
        run = paragraph.add_run()
        font = run.font
        if fmt.bold:
            font.bold = True
        if fmt.italic:
            font.italic = True
        if fmt.underline or fmt.href:
            font.underline = True
        if fmt.strike:
            font.strike = True
        if fmt.superscript:
            font.superscript = True
        if fmt.subscript:
            font.subscript = True
        color = fmt.color or ("0563C1" if fmt.href else None)
        if color:
            font.color.rgb = RGBColor.from_string(color)
        family = "Courier New" if fmt.code and not fmt.font else fmt.font
        if family:
            font.name = family
            self._set_east_asian_font(run._r, family)
        if fmt.size_pt:
            font.size = Pt(round(fmt.size_pt * 2) / 2)
        if fmt.background:
            highlight = HIGHLIGHT_BY_HEX.get(f"#{fmt.background.lower()}")
            r_pr = run._r.get_or_add_rPr()
            if highlight:
                element = OxmlElement("w:highlight")
                element.set(qn("w:val"), highlight)
            else:
                element = OxmlElement("w:shd")
                element.set(qn("w:val"), "clear")
                element.set(qn("w:color"), "auto")
                element.set(qn("w:fill"), fmt.background)
            r_pr.append(element)

        if fmt.href:
            self._wrap_hyperlink(paragraph, run, fmt.href)
        if fmt.revision:
            self._wrap_revision(run, fmt.revision)
        for comment_id, author, date, text in fmt.comments:
            anchor = self._comment_anchors.setdefault(comment_id or f"c{len(self._comment_anchors)}", _CommentAnchor(author, date, text))
            anchor.runs.append(run)
        return run

    def _wrap_hyperlink(self, paragraph: Any, run: Any, href: str) -> None:
        run_element = run._r
        previous = run_element.getprevious()
        # Merge consecutive runs of the same link into one w:hyperlink.
        if self._last_link and previous is self._last_link[0] and self._last_link[1] == href:
            previous.append(run_element)
            return
        hyperlink = OxmlElement("w:hyperlink")
        if href.startswith("#"):
            hyperlink.set(qn("w:anchor"), href[1:])
        else:
            relationship_id = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
            hyperlink.set(qn("r:id"), relationship_id)
        run_element.addprevious(hyperlink)
        hyperlink.append(run_element)
        self._last_link = (hyperlink, href)

    def _wrap_revision(self, run: Any, revision: tuple[str, str, str]) -> None:
        kind, author, date = revision
        self._revision_id += 1
        wrapper = OxmlElement("w:del" if kind == "delete" else "w:ins")
        wrapper.set(qn("w:id"), str(self._revision_id))
        wrapper.set(qn("w:author"), author or "GolinelliAI")
        wrapper.set(qn("w:date"), date or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
        element = run._r
        parent = element.getparent()
        # Keep hyperlinks intact: wrap the run where it sits.
        element.addprevious(wrapper)
        wrapper.append(element)
        del parent

    # images -------------------------------------------------------------------------------
    def _write_image(self, node: Tag, paragraph: Any) -> None:
        source = str(node.get("src") or "")
        match = re.match(r"data:([^;]+);base64,(.+)", source, re.DOTALL)
        if not match:
            alt = str(node.get("alt") or "")
            if alt:
                paragraph.add_run(f"[{alt}]").italic = True
            return
        try:
            data = base64.b64decode(match.group(2))
            with Image.open(io.BytesIO(data)) as image:
                natural_width, natural_height = image.size
                if image.format not in {"PNG", "JPEG", "GIF", "BMP", "TIFF"}:
                    converted = io.BytesIO()
                    image.convert("RGBA").save(converted, format="PNG")
                    data = converted.getvalue()
        except Exception:
            return
        style = parse_style(node.get("style"))
        width_px = parse_length_px(str(node.get("width") or "")) or parse_length_px(style.get("width")) or natural_width
        height_px = parse_length_px(str(node.get("height") or "")) or parse_length_px(style.get("height"))
        if not height_px and natural_width:
            height_px = width_px * natural_height / natural_width
        max_width = self.text_width_px
        if width_px > max_width:
            scale = max_width / width_px
            width_px, height_px = max_width, (height_px or 0) * scale
        run = paragraph.add_run()
        kwargs: dict[str, Any] = {"width": Inches(width_px / PX_PER_INCH)}
        if height_px:
            kwargs["height"] = Inches(height_px / PX_PER_INCH)
        try:
            run.add_picture(io.BytesIO(data), **kwargs)
        except Exception:
            paragraph._p.remove(run._r)

    # comments -----------------------------------------------------------------------------
    def _flush_comments(self) -> None:
        for anchor in self._comment_anchors.values():
            runs = [run for run in anchor.runs if run._r.getparent() is not None]
            if not runs or not anchor.text:
                continue
            # python-docx anchors a comment on runs of a single paragraph: use the first paragraph span.
            first_paragraph = self._paragraph_of(runs[0])
            same_paragraph = [run for run in runs if self._paragraph_of(run) is first_paragraph]
            try:
                initials = "".join(part[:1] for part in anchor.author.split()[:2]).upper() or None
                self.document.add_comment(
                    same_paragraph if len(same_paragraph) > 1 else same_paragraph[0],
                    text=anchor.text,
                    author=anchor.author or "GolinelliAI",
                    initials=initials,
                )
            except Exception:
                continue

    @staticmethod
    def _paragraph_of(run: Any) -> Any:
        element = run._r
        while element is not None and element.tag != qn("w:p"):
            element = element.getparent()
        return element


def html_to_docx(
    raw_html: str,
    *,
    title: str = "",
    margins_px: dict[str, Any] | None = None,
    header: dict[str, Any] | None = None,
    page: dict[str, Any] | None = None,
    header_footer: dict[str, Any] | None = None,
    page_breaks: Iterable[int] | None = None,
) -> bytes:
    writer = DocxWriter(title=title, margins_px=margins_px, page=page)
    writer.write_header_card(header)
    writer.write_html(raw_html, page_breaks)
    writer.write_header_footer(header_footer, raw_html)
    body = writer.document.element.body
    blocks = [child for child in body if child.tag != qn("w:sectPr")]
    if not blocks or blocks[-1].tag == qn("w:tbl"):
        # Word expects the body to end with a paragraph.
        writer.document.add_paragraph()
    return writer.save()


__all__ = ["html_to_docx", "DocxWriter", "parse_color", "parse_style"]
