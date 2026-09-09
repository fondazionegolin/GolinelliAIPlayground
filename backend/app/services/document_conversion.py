"""Import/export bridge between LMS files and the native document editors.

The original upload is always kept separately in object storage.  This module only
creates an editable representation or a newly exported copy, so conversions are
non-destructive by design.
"""

from __future__ import annotations

import base64
import csv
import html
import io
import json
import mimetypes
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import fitz
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches as PptxInches, Pt as PptxPt


SUPPORTED_IMPORT_EXTENSIONS = {"pdf", "ppt", "pptx", "doc", "docx", "md", "xls", "xlsx", "csv"}
SUPPORTED_EXPORT_FORMATS = {"pdf", "ppt", "pptx", "doc", "docx", "xlsx"}

MIME_BY_EXTENSION = {
    "pdf": "application/pdf",
    "ppt": "application/vnd.ms-powerpoint",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "doc": "application/msword",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "md": "text/markdown",
    "xls": "application/vnd.ms-excel",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
}


@dataclass(slots=True)
class ImportedDocument:
    title: str
    doc_type: str
    content_json: str
    extension: str


@dataclass(slots=True)
class ExportedDocument:
    filename: str
    mime_type: str
    content: bytes


def normalized_extension(filename: str) -> str:
    return Path(filename or "").suffix.lower().lstrip(".")


def _safe_title(filename: str) -> str:
    title = Path(filename or "Documento").stem.strip()
    return title[:255] or "Documento"


def _source_metadata(filename: str, mime_type: str, file_id: str | None) -> dict[str, Any]:
    extension = normalized_extension(filename)
    return {
        "filename": Path(filename).name,
        "extension": extension,
        "mimeType": mime_type or MIME_BY_EXTENSION.get(extension, "application/octet-stream"),
        "fileId": file_id,
        "url": f"/api/v1/files/{file_id}/content" if file_id else None,
        "preservedOriginal": True,
    }


def _run_libreoffice(input_bytes: bytes, source_ext: str, target_ext: str) -> bytes:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice non è disponibile per questa conversione")
    with tempfile.TemporaryDirectory(prefix="eduai-office-") as temp_dir:
        root = Path(temp_dir)
        source = root / f"source.{source_ext}"
        source.write_bytes(input_bytes)
        command = [
            executable,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            "--convert-to",
            target_ext,
            "--outdir",
            str(root),
            str(source),
        ]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
        output = root / f"source.{target_ext}"
        if completed.returncode != 0 or not output.exists():
            detail = (completed.stderr or completed.stdout or "conversione non riuscita").strip()
            raise RuntimeError(f"LibreOffice: {detail[:400]}")
        return output.read_bytes()


def _markdown_to_html(raw: str) -> str:
    lines = raw.replace("\r\n", "\n").split("\n")
    rendered: list[str] = []
    in_list = False
    for line in lines:
        stripped = line.strip()
        if stripped.startswith(("- ", "* ")):
            if not in_list:
                rendered.append("<ul>")
                in_list = True
            rendered.append(f"<li>{html.escape(stripped[2:])}</li>")
            continue
        if in_list:
            rendered.append("</ul>")
            in_list = False
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            rendered.append(f"<h{level}>{html.escape(heading.group(2))}</h{level}>")
        elif stripped:
            value = html.escape(stripped)
            value = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", value)
            value = re.sub(r"`(.+?)`", r"<code>\1</code>", value)
            rendered.append(f"<p>{value}</p>")
        else:
            rendered.append("<p></p>")
    if in_list:
        rendered.append("</ul>")
    return "\n".join(rendered) or "<p></p>"


_DOCX_HIGHLIGHT_COLORS = {
    "black": "#000000",
    "blue": "#0000ff",
    "cyan": "#00ffff",
    "darkBlue": "#000080",
    "darkCyan": "#008080",
    "darkGray": "#808080",
    "darkGreen": "#008000",
    "darkMagenta": "#800080",
    "darkRed": "#800000",
    "darkYellow": "#808000",
    "green": "#00ff00",
    "lightGray": "#c0c0c0",
    "magenta": "#ff00ff",
    "red": "#ff0000",
    "white": "#ffffff",
    "yellow": "#ffff00",
}
_VML_IMAGE_TAG = "{urn:schemas-microsoft-com:vml}imagedata"


def _docx_bool(element: Any | None) -> bool:
    if element is None:
        return False
    return str(element.get(qn("w:val"), "true")).lower() not in {"0", "false", "off", "none"}


def _docx_comments(data: bytes) -> dict[str, dict[str, str]]:
    comments: dict[str, dict[str, str]] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            raw = archive.read("word/comments.xml")
    except (KeyError, zipfile.BadZipFile):
        return comments

    # python-docx already depends on lxml and returns the same OOXML element type.
    from docx.oxml import parse_xml

    root = parse_xml(raw)
    for node in root:
        if node.tag != qn("w:comment"):
            continue
        comment_id = str(node.get(qn("w:id"), ""))
        text = "".join(
            str(item.text or "")
            for item in node.iter()
            if item.tag in {qn("w:t"), qn("w:delText")}
        ).strip()
        comments[comment_id] = {
            "author": str(node.get(qn("w:author"), "")),
            "date": str(node.get(qn("w:date"), "")),
            "text": text,
        }
    return comments


def _docx_style_context(document: Any) -> dict[str, Any]:
    names: dict[str, str] = {}
    elements: dict[str, Any] = {}
    for style in document.styles.element:
        if style.tag != qn("w:style"):
            continue
        style_id = str(style.get(qn("w:styleId"), ""))
        name = style.find(qn("w:name"))
        names[style_id] = str(name.get(qn("w:val"), style_id)) if name is not None else style_id
        elements[style_id] = style
    defaults = document.styles.element.find(qn("w:docDefaults"))
    run_defaults = None
    if defaults is not None:
        run_properties_default = defaults.find(qn("w:rPrDefault"))
        if run_properties_default is not None:
            run_defaults = run_properties_default.find(qn("w:rPr"))
    return {"names": names, "elements": elements, "run_defaults": run_defaults}


def _docx_style_chain(style_id: str, style_context: dict[str, Any]) -> list[Any]:
    elements = style_context["elements"]
    chain: list[Any] = []
    seen: set[str] = set()
    current_id = style_id
    while current_id and current_id not in seen:
        seen.add(current_id)
        style = elements.get(current_id)
        if style is None:
            break
        chain.append(style)
        based_on = style.find(qn("w:basedOn"))
        current_id = str(based_on.get(qn("w:val"), "")) if based_on is not None else ""
    chain.reverse()
    return chain


def _docx_run_property(
    run_properties: Any | None,
    paragraph_style_id: str,
    style_context: dict[str, Any],
    property_name: str,
) -> Any | None:
    sources: list[Any] = []
    if style_context.get("run_defaults") is not None:
        sources.append(style_context["run_defaults"])
    for style in _docx_style_chain(paragraph_style_id, style_context):
        inherited = style.find(qn("w:rPr"))
        if inherited is not None:
            sources.append(inherited)
    if run_properties is not None:
        character_style = run_properties.find(qn("w:rStyle"))
        character_style_id = (
            str(character_style.get(qn("w:val"), "")) if character_style is not None else ""
        )
        for style in _docx_style_chain(character_style_id, style_context):
            inherited = style.find(qn("w:rPr"))
            if inherited is not None:
                sources.append(inherited)
        sources.append(run_properties)
    result = None
    for source in sources:
        candidate = source.find(qn(f"w:{property_name}"))
        if candidate is not None:
            result = candidate
    return result


def _docx_relationship(part: Any, relationship_id: str) -> Any | None:
    try:
        return part.rels[relationship_id]
    except (KeyError, TypeError):
        return None


def _docx_image_html(element: Any, part: Any) -> str:
    relationship_id = ""
    for descendant in element.iter():
        if descendant.tag == qn("a:blip"):
            relationship_id = str(descendant.get(qn("r:embed"), ""))
            break
        if descendant.tag == _VML_IMAGE_TAG:
            relationship_id = str(descendant.get(qn("r:id"), ""))
            break
    relationship = _docx_relationship(part, relationship_id)
    if relationship is None:
        return ""
    target_part = getattr(relationship, "target_part", None)
    blob = getattr(target_part, "blob", None)
    if not blob:
        return ""

    alt = "Immagine importata"
    width = ""
    height = ""
    for descendant in element.iter():
        if descendant.tag == qn("wp:docPr"):
            alt = str(
                descendant.get("descr")
                or descendant.get("title")
                or descendant.get("name")
                or alt
            )
        elif descendant.tag == qn("wp:extent"):
            try:
                width = str(max(1, round(int(descendant.get("cx", "0")) / 9525)))
                height = str(max(1, round(int(descendant.get("cy", "0")) / 9525)))
            except (TypeError, ValueError):
                width = height = ""

    encoded = base64.b64encode(blob).decode("ascii")
    content_type = str(getattr(target_part, "content_type", "image/png"))
    dimensions = f' width="{width}" height="{height}"' if width and height else ""
    return (
        f'<img src="data:{html.escape(content_type, quote=True)};base64,{encoded}"'
        f' alt="{html.escape(alt, quote=True)}"{dimensions}>'
    )


def _docx_run_html(
    run: Any,
    part: Any,
    paragraph_style_id: str,
    style_context: dict[str, Any],
) -> str:
    fragments: list[str] = []
    for child in run:
        if child.tag in {qn("w:t"), qn("w:delText"), qn("w:instrText")}:
            fragments.append(html.escape(str(child.text or "")))
        elif child.tag == qn("w:tab"):
            fragments.append("&emsp;")
        elif child.tag in {qn("w:br"), qn("w:cr")}:
            fragments.append("<br>")
        elif child.tag == qn("w:noBreakHyphen"):
            fragments.append("&#8209;")
        elif child.tag in {qn("w:drawing"), qn("w:pict"), qn("w:object")}:
            fragments.append(_docx_image_html(child, part))
    content = "".join(fragments)
    if not content:
        return ""

    properties = run.find(qn("w:rPr"))
    wrappers: list[str] = []
    styles: list[str] = []
    if _docx_bool(_docx_run_property(properties, paragraph_style_id, style_context, "b")):
        wrappers.append("strong")
    if _docx_bool(_docx_run_property(properties, paragraph_style_id, style_context, "i")):
        wrappers.append("em")
    if _docx_bool(_docx_run_property(properties, paragraph_style_id, style_context, "u")):
        wrappers.append("u")
    if _docx_bool(_docx_run_property(properties, paragraph_style_id, style_context, "strike")) or _docx_bool(
        _docx_run_property(properties, paragraph_style_id, style_context, "dstrike")
    ):
        wrappers.append("s")

    vertical = _docx_run_property(properties, paragraph_style_id, style_context, "vertAlign")
    vertical_value = str(vertical.get(qn("w:val"), "")) if vertical is not None else ""
    if vertical_value == "superscript":
        wrappers.append("sup")
    elif vertical_value == "subscript":
        wrappers.append("sub")

    color = _docx_run_property(properties, paragraph_style_id, style_context, "color")
    color_value = str(color.get(qn("w:val"), "")) if color is not None else ""
    if re.fullmatch(r"[0-9A-Fa-f]{6}", color_value):
        styles.append(f"color:#{color_value}")

    highlight = _docx_run_property(properties, paragraph_style_id, style_context, "highlight")
    highlight_value = str(highlight.get(qn("w:val"), "")) if highlight is not None else ""
    if highlight_value in _DOCX_HIGHLIGHT_COLORS:
        styles.append(f"background-color:{_DOCX_HIGHLIGHT_COLORS[highlight_value]}")

    shading = _docx_run_property(properties, paragraph_style_id, style_context, "shd")
    shading_fill = str(shading.get(qn("w:fill"), "")) if shading is not None else ""
    if re.fullmatch(r"[0-9A-Fa-f]{6}", shading_fill) and shading_fill.lower() != "auto":
        styles.append(f"background-color:#{shading_fill}")

    size = _docx_run_property(properties, paragraph_style_id, style_context, "sz")
    try:
        half_points = int(size.get(qn("w:val"), "0")) if size is not None else 0
    except (TypeError, ValueError):
        half_points = 0
    if half_points:
        styles.append(f"font-size:{half_points / 2:g}pt")

    fonts = _docx_run_property(properties, paragraph_style_id, style_context, "rFonts")
    if fonts is not None:
        family = (
            fonts.get(qn("w:ascii"))
            or fonts.get(qn("w:hAnsi"))
            or fonts.get(qn("w:eastAsia"))
        )
        if family:
            safe_family = str(family).replace('"', "").replace("'", "")
            styles.append(f"font-family:'{safe_family}'")

    if styles:
        content = f'<span style="{html.escape(";".join(styles), quote=True)}">{content}</span>'
    for wrapper in wrappers:
        content = f"<{wrapper}>{content}</{wrapper}>"
    return content


def _docx_annotation_attributes(metadata: dict[str, str], prefix: str) -> str:
    return " ".join(
        f'data-{prefix}-{name}="{html.escape(str(value), quote=True)}"'
        for name, value in metadata.items()
        if value
    )


def _docx_inline_html(
    parent: Any,
    part: Any,
    comments: dict[str, dict[str, str]],
    active_comments: set[str],
    paragraph_style_id: str,
    style_context: dict[str, Any],
) -> str:
    output: list[str] = []
    for child in parent:
        if child.tag in {qn("w:pPr"), qn("w:rPr")}:
            continue
        if child.tag == qn("w:commentRangeStart"):
            active_comments.add(str(child.get(qn("w:id"), "")))
            continue
        if child.tag == qn("w:commentRangeEnd"):
            active_comments.discard(str(child.get(qn("w:id"), "")))
            continue

        fragment = ""
        if child.tag == qn("w:r"):
            fragment = _docx_run_html(child, part, paragraph_style_id, style_context)
            references = [
                str(reference.get(qn("w:id"), ""))
                for reference in child.iter()
                if reference.tag == qn("w:commentReference")
            ]
            if not fragment and references:
                fragment = "&#8203;"
        elif child.tag == qn("w:hyperlink"):
            fragment = _docx_inline_html(
                child,
                part,
                comments,
                active_comments,
                paragraph_style_id,
                style_context,
            )
            relationship_id = str(child.get(qn("r:id"), ""))
            relationship = _docx_relationship(part, relationship_id)
            href = str(getattr(relationship, "target_ref", "") or "")
            anchor = str(child.get(qn("w:anchor"), "") or "")
            if anchor and not href:
                href = f"#{anchor}"
            if href:
                fragment = f'<a href="{html.escape(href, quote=True)}">{fragment}</a>'
        elif child.tag in {qn("w:ins"), qn("w:del")}:
            fragment = _docx_inline_html(
                child,
                part,
                comments,
                active_comments,
                paragraph_style_id,
                style_context,
            )
            kind = "insert" if child.tag == qn("w:ins") else "delete"
            metadata = {
                "author": str(child.get(qn("w:author"), "")),
                "date": str(child.get(qn("w:date"), "")),
            }
            attributes = _docx_annotation_attributes(metadata, "revision")
            fragment = f'<span data-revision-kind="{kind}" {attributes}>{fragment}</span>'
        elif child.tag in {qn("w:smartTag"), qn("w:sdt"), qn("w:sdtContent"), qn("w:fldSimple")}:
            fragment = _docx_inline_html(
                child,
                part,
                comments,
                active_comments,
                paragraph_style_id,
                style_context,
            )

        if fragment and active_comments:
            for comment_id in sorted(active_comments):
                metadata = comments.get(comment_id, {})
                attributes = _docx_annotation_attributes(
                    {"id": comment_id, **metadata},
                    "comment",
                )
                fragment = f'<span data-docx-comment="true" {attributes}>{fragment}</span>'
        output.append(fragment)
    return "".join(output)


def _docx_paragraph_html(
    paragraph: Any,
    part: Any,
    comments: dict[str, dict[str, str]],
    active_comments: set[str],
    style_context: dict[str, Any],
) -> str:
    properties = paragraph.find(qn("w:pPr"))
    style_id = ""
    alignment = ""
    is_list = False
    if properties is not None:
        style = properties.find(qn("w:pStyle"))
        style_id = str(style.get(qn("w:val"), "")) if style is not None else ""
        justification = properties.find(qn("w:jc"))
        alignment_value = str(justification.get(qn("w:val"), "")) if justification is not None else ""
        alignment = {
            "both": "justify",
            "distribute": "justify",
            "center": "center",
            "right": "right",
            "left": "left",
        }.get(alignment_value, "")
        is_list = properties.find(qn("w:numPr")) is not None

    style_name = style_context["names"].get(style_id, style_id)
    heading = re.search(r"(?:heading|titolo)\s*([1-6])", style_name, re.IGNORECASE)
    content = _docx_inline_html(
        paragraph,
        part,
        comments,
        active_comments,
        style_id,
        style_context,
    ) or "<br>"
    style_attribute = f' style="text-align:{alignment}"' if alignment else ""
    if heading:
        tag = f"h{heading.group(1)}"
        return f"<{tag}{style_attribute}>{content}</{tag}>"
    if is_list or "list" in style_name.lower() or "elenco" in style_name.lower():
        return f"<ul><li>{content}</li></ul>"
    return f"<p{style_attribute}>{content}</p>"


def _docx_table_html(
    table: Any,
    part: Any,
    comments: dict[str, dict[str, str]],
    style_context: dict[str, Any],
) -> str:
    rows: list[str] = []
    active_comments: set[str] = set()
    row_elements = [child for child in table if child.tag == qn("w:tr")]
    for row_index, row in enumerate(row_elements):
        cells: list[str] = []
        for cell in (child for child in row if child.tag == qn("w:tc")):
            properties = cell.find(qn("w:tcPr"))
            attributes: list[str] = []
            styles: list[str] = []
            if properties is not None:
                grid_span = properties.find(qn("w:gridSpan"))
                if grid_span is not None and grid_span.get(qn("w:val")):
                    attributes.append(f'colspan="{html.escape(str(grid_span.get(qn("w:val"))), quote=True)}"')
                shading = properties.find(qn("w:shd"))
                fill = str(shading.get(qn("w:fill"), "")) if shading is not None else ""
                if re.fullmatch(r"[0-9A-Fa-f]{6}", fill) and fill.lower() != "auto":
                    styles.append(f"background-color:#{fill}")
            if styles:
                attributes.append(f'style="{html.escape(";".join(styles), quote=True)}"')
            blocks = [
                _docx_paragraph_html(child, part, comments, active_comments, style_context)
                for child in cell
                if child.tag == qn("w:p")
            ]
            tag = "th" if row_index == 0 else "td"
            cells.append(f"<{tag} {' '.join(attributes)}>{''.join(blocks) or '<p><br></p>'}</{tag}>")
        rows.append(f"<tr>{''.join(cells)}</tr>")
    return f"<table><tbody>{''.join(rows)}</tbody></table>"


def _docx_part_blocks(
    container: Any,
    part: Any,
    comments: dict[str, dict[str, str]],
    style_context: dict[str, Any],
) -> list[str]:
    output: list[str] = []
    active_comments: set[str] = set()
    for child in container:
        if child.tag == qn("w:p"):
            output.append(_docx_paragraph_html(child, part, comments, active_comments, style_context))
        elif child.tag == qn("w:tbl"):
            output.append(_docx_table_html(child, part, comments, style_context))
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                output.extend(_docx_part_blocks(content, part, comments, style_context))
    return output


def _docx_region_html(kind: str, label: str, blocks: list[str]) -> str:
    meaningful = "".join(blocks).replace("<p><br></p>", "").strip()
    if not meaningful:
        return ""
    return (
        f'<section data-docx-region="{kind}" data-docx-region-label="{html.escape(label, quote=True)}">'
        f"{''.join(blocks)}</section>"
    )


def _docx_to_html(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    comments = _docx_comments(data)
    style_context = _docx_style_context(document)
    output: list[str] = []
    seen_parts: set[str] = set()

    for section in document.sections:
        for kind, label, region in (
            ("header", "Intestazione", section.header),
            ("header-first", "Intestazione prima pagina", section.first_page_header),
            ("header-even", "Intestazione pagine pari", section.even_page_header),
        ):
            part_name = str(region.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            output.append(
                _docx_region_html(
                    kind,
                    label,
                    _docx_part_blocks(region._element, region.part, comments, style_context),
                )
            )

    output.extend(_docx_part_blocks(document.element.body, document.part, comments, style_context))

    for section in document.sections:
        for kind, label, region in (
            ("footer", "Piè di pagina", section.footer),
            ("footer-first", "Piè di pagina prima pagina", section.first_page_footer),
            ("footer-even", "Piè di pagina pagine pari", section.even_page_footer),
        ):
            part_name = str(region.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            output.append(
                _docx_region_html(
                    kind,
                    label,
                    _docx_part_blocks(region._element, region.part, comments, style_context),
                )
            )
    return "\n".join(fragment for fragment in output if fragment) or "<p></p>"


def _pdf_to_html(data: bytes) -> str:
    pdf = fitz.open(stream=data, filetype="pdf")
    pages: list[str] = []
    try:
        for page_index, page in enumerate(pdf):
            blocks = sorted(page.get_text("blocks"), key=lambda block: (block[1], block[0]))
            paragraphs = [f"<p>{html.escape(str(block[4]).strip()).replace(chr(10), '<br>')}</p>" for block in blocks if str(block[4]).strip()]
            pages.append(
                f'<section data-imported-page="{page_index + 1}" style="break-after:page">'
                f"{''.join(paragraphs) or '<p></p>'}</section>"
            )
    finally:
        pdf.close()
    return "\n".join(pages) or "<p></p>"


def _pdf_first_page_preview(data: bytes) -> str | None:
    pdf = fitz.open(stream=data, filetype="pdf")
    try:
        if not pdf.page_count:
            return None
        page = pdf.load_page(0)
        scale = min(2.0, 1200 / max(1, page.rect.width))
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        encoded = base64.b64encode(pixmap.tobytes("jpeg", jpg_quality=82)).decode("ascii")
        return f"data:image/jpeg;base64,{encoded}"
    finally:
        pdf.close()


def _pptx_to_native(data: bytes) -> dict[str, Any]:
    presentation = Presentation(io.BytesIO(data))
    width = int(presentation.slide_width or 1)
    height = int(presentation.slide_height or 1)
    format_name = "4:3" if abs((width / height) - (4 / 3)) < 0.08 else "16:9"
    target_width, target_height = ((800, 600) if format_name == "4:3" else (960, 540))
    slides: list[dict[str, Any]] = []
    for slide_index, slide in enumerate(presentation.slides):
        blocks: list[dict[str, Any]] = []
        title = f"Slide {slide_index + 1}"
        for shape in slide.shapes:
            x = round(shape.left / width * target_width, 2)
            y = round(shape.top / height * target_height, 2)
            block_width = max(8, round(shape.width / width * target_width, 2))
            block_height = max(8, round(shape.height / height * target_height, 2))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                encoded = base64.b64encode(image.blob).decode("ascii")
                blocks.append({
                    "id": str(uuid.uuid4()), "type": "image", "content": f"data:{image.content_type};base64,{encoded}",
                    "x": x, "y": y, "width": block_width, "height": block_height, "style": {},
                })
                continue
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            text = shape.text.strip()
            if getattr(slide.shapes, "title", None) is shape:
                title = text.splitlines()[0][:200]
            paragraph = shape.text_frame.paragraphs[0] if shape.text_frame.paragraphs else None
            run = paragraph.runs[0] if paragraph and paragraph.runs else None
            font_size = round((run.font.size.pt if run and run.font.size else (30 if y < 120 else 18)) * target_width / 960, 1)
            color = "#1e293b"
            try:
                if run and run.font.color and run.font.color.type and run.font.color.rgb:
                    color = f"#{run.font.color.rgb}"
            except AttributeError:
                pass
            alignment = "left"
            if paragraph and paragraph.alignment == PP_ALIGN.CENTER:
                alignment = "center"
            elif paragraph and paragraph.alignment == PP_ALIGN.RIGHT:
                alignment = "right"
            blocks.append({
                "id": str(uuid.uuid4()), "type": "text", "content": text,
                "x": x, "y": y, "width": block_width, "height": block_height,
                "style": {"fontSize": font_size, "color": color, "fontWeight": "bold" if run and run.font.bold else "normal", "textAlign": alignment, "lineHeight": 1.2},
            })
        slides.append({"id": str(uuid.uuid4()), "title": title, "blocks": blocks})
    return {"type": "presentation_v2", "format": format_name, "slides": slides or [{"id": str(uuid.uuid4()), "title": "Slide 1", "blocks": []}]}


def _xlsx_to_native(data: bytes) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    sheet = workbook.active
    rows: list[list[str]] = []
    for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
        if row_index >= 1000:
            break
        values = ["" if value is None else str(value) for value in row[:100]]
        while values and values[-1] == "":
            values.pop()
        rows.append(values)
    while rows and not rows[-1]:
        rows.pop()
    return {"type": "sheet_v1", "data": rows or [[""]], "sheetName": sheet.title}


def _csv_to_native(data: bytes) -> dict[str, Any]:
    raw = data.decode("utf-8-sig", errors="replace")
    sample = raw[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[str]] = []
    for row_index, row in enumerate(csv.reader(io.StringIO(raw), dialect)):
        if row_index >= 1000:
            break
        values = [str(value) for value in row[:100]]
        while values and values[-1] == "":
            values.pop()
        rows.append(values)
    while rows and not rows[-1]:
        rows.pop()
    return {"type": "sheet_v1", "data": rows or [[""]], "sheetName": "CSV"}


def import_document(filename: str, data: bytes, mime_type: str = "", file_id: str | None = None) -> ImportedDocument:
    extension = normalized_extension(filename)
    if extension not in SUPPORTED_IMPORT_EXTENSIONS:
        raise ValueError(f"Formato .{extension or '?'} non supportato")
    working = data
    modern_extension = extension
    if extension in {"doc", "ppt", "xls"}:
        modern_extension = {"doc": "docx", "ppt": "pptx", "xls": "xlsx"}[extension]
        working = _run_libreoffice(data, extension, modern_extension)
    source = _source_metadata(filename, mime_type, file_id)
    if modern_extension == "pptx":
        content = _pptx_to_native(working)
        doc_type = "presentation"
    elif modern_extension == "xlsx":
        content = _xlsx_to_native(working)
        doc_type = "sheet"
    elif extension == "csv":
        content = _csv_to_native(data)
        doc_type = "sheet"
    elif modern_extension == "docx":
        content = {"type": "document_v1", "htmlContent": _docx_to_html(working)}
        doc_type = "document"
    elif extension == "md":
        content = {"type": "document_v1", "htmlContent": _markdown_to_html(data.decode("utf-8", errors="replace"))}
        doc_type = "document"
    else:
        content = {
            "type": "document_v1",
            "htmlContent": _pdf_to_html(data),
            "previewImage": _pdf_first_page_preview(data),
        }
        doc_type = "document"
    content["source"] = source
    content["imported"] = True
    return ImportedDocument(_safe_title(filename), doc_type, json.dumps(content, ensure_ascii=False), extension)


def _append_html_to_docx(document: Document, raw_html: str) -> None:
    soup = BeautifulSoup(raw_html or "<p></p>", "html.parser")
    for node in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table", "img"]):
        if node.find_parent(["p", "li", "table"]) and node.name != "img":
            continue
        if node.name.startswith("h"):
            document.add_heading(node.get_text(" ", strip=True), level=min(6, int(node.name[1])))
        elif node.name == "li":
            document.add_paragraph(node.get_text(" ", strip=True), style="List Bullet")
        elif node.name == "table":
            rows = node.find_all("tr")
            column_count = max((len(row.find_all(["th", "td"])) for row in rows), default=1)
            table = document.add_table(rows=max(1, len(rows)), cols=max(1, column_count))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for column_index, cell in enumerate(row.find_all(["th", "td"])):
                    table.cell(row_index, column_index).text = cell.get_text(" ", strip=True)
        elif node.name == "img":
            source = str(node.get("src") or "")
            match = re.match(r"data:([^;]+);base64,(.+)", source, re.DOTALL)
            if match:
                try:
                    document.add_picture(io.BytesIO(base64.b64decode(match.group(2))), width=Inches(5.8))
                except Exception:
                    pass
        else:
            paragraph = document.add_paragraph(node.get_text(" ", strip=True))
            style = str(node.get("style") or "")
            if "text-align: center" in style:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _native_to_docx(content: dict[str, Any], title: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    document.add_heading(title, level=0)
    if content.get("type") == "presentation_v2" or isinstance(content.get("slides"), list):
        for index, slide in enumerate(content.get("slides") or []):
            document.add_heading(slide.get("title") or f"Slide {index + 1}", level=1)
            for block in sorted(slide.get("blocks") or [], key=lambda item: (item.get("y", 0), item.get("x", 0))):
                if block.get("type") == "text" and block.get("content"):
                    document.add_paragraph(str(block["content"]))
                elif block.get("type") == "image" and str(block.get("content", "")).startswith("data:"):
                    match = re.match(r"data:([^;]+);base64,(.+)", block["content"], re.DOTALL)
                    if match:
                        try:
                            document.add_picture(io.BytesIO(base64.b64decode(match.group(2))), width=Inches(5.8))
                        except Exception:
                            pass
    elif content.get("type") == "sheet_v1" or isinstance(content.get("data"), list):
        rows = content.get("data") or []
        columns = max((len(row) for row in rows), default=1)
        table = document.add_table(rows=max(1, len(rows)), cols=max(1, columns))
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = str(value)
    else:
        _append_html_to_docx(document, str(content.get("htmlContent") or content.get("content") or ""))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx_rgb(value: Any, fallback: str = "000000") -> PptxRGBColor:
    raw = str(value or "").strip().lstrip("#")
    if len(raw) == 3 and re.fullmatch(r"[0-9a-fA-F]{3}", raw):
        raw = "".join(character * 2 for character in raw)
    if not re.fullmatch(r"[0-9a-fA-F]{6}", raw):
        raw = fallback
    return PptxRGBColor.from_string(raw.upper())


def _is_transparent(value: Any) -> bool:
    return not value or str(value).strip().lower() in {"transparent", "none"}


def _pptx_box(block: dict[str, Any], canvas_width: float, canvas_height: float) -> tuple[float, float, float, float]:
    x = max(0.0, min(canvas_width, float(block.get("x", 0) or 0)))
    y = max(0.0, min(canvas_height, float(block.get("y", 0) or 0)))
    width = max(1.0, float(block.get("width", 1) or 1))
    height = max(1.0, float(block.get("height", 1) or 1))
    return x, y, min(width, max(1.0, canvas_width - x)), min(height, max(1.0, canvas_height - y))


def _set_shape_rotation(shape: Any, block: dict[str, Any]) -> None:
    try:
        shape.rotation = float(block.get("rotation") or 0) % 360
    except (TypeError, ValueError):
        shape.rotation = 0


def _add_text_to_slide(
    slide: Any,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    style: dict[str, Any],
    scale_x: float,
    scale_y: float,
    rotation: float = 0,
) -> None:
    box = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(width), PptxInches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    padding = max(0.0, float(style.get("padding") or 0))
    frame.margin_left = frame.margin_right = PptxInches(padding * scale_x)
    frame.margin_top = frame.margin_bottom = PptxInches(padding * scale_y)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    run = paragraph.runs[0]
    # The editor stores CSS pixels; Office uses points (1 CSS px = 0.75 pt).
    run.font.size = PptxPt(max(6, min(72, float(style.get("fontSize") or 18) * 0.75)))
    if style.get("fontFamily"):
        run.font.name = str(style["fontFamily"])
    run.font.color.rgb = _pptx_rgb(style.get("color"), "000000")
    run.font.bold = str(style.get("fontWeight") or "").lower() in {"bold", "600", "700", "800", "900"}
    run.font.italic = str(style.get("fontStyle") or "").lower() == "italic"
    run.font.underline = "underline" in str(style.get("textDecoration") or "").lower()
    alignment = style.get("textAlign")
    paragraph.alignment = (
        PP_ALIGN.CENTER if alignment == "center"
        else PP_ALIGN.RIGHT if alignment == "right"
        else PP_ALIGN.JUSTIFY if alignment == "justify"
        else PP_ALIGN.LEFT
    )
    line_height = style.get("lineHeight")
    if isinstance(line_height, (int, float)) and line_height > 0:
        paragraph.line_spacing = float(line_height)
    background = style.get("backgroundColor")
    if not _is_transparent(background):
        box.fill.solid()
        box.fill.fore_color.rgb = _pptx_rgb(background, "FFFFFF")
    else:
        box.fill.background()
    box.line.fill.background()
    box.rotation = rotation % 360


def _cover_image_bytes(data: bytes, width: float, height: float) -> bytes:
    with Image.open(io.BytesIO(data)) as source:
        target_width = max(1, min(3840, round(width * 2)))
        target_height = max(1, min(2160, round(height * 2)))
        image = ImageOps.fit(source.convert("RGBA"), (target_width, target_height), method=Image.Resampling.LANCZOS)
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()


def _add_shape_to_slide(
    slide: Any,
    block: dict[str, Any],
    x: float,
    y: float,
    width: float,
    height: float,
) -> None:
    style = block.get("style") or {}
    if block.get("type") == "line":
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            PptxInches(x),
            PptxInches(y),
            PptxInches(x + width),
            PptxInches(y + height),
        )
    else:
        shape_type = MSO_SHAPE.OVAL if block.get("type") == "ellipse" else (
            MSO_SHAPE.ROUNDED_RECTANGLE if float(style.get("cornerRadius") or 0) > 0 else MSO_SHAPE.RECTANGLE
        )
        shape = slide.shapes.add_shape(
            shape_type,
            PptxInches(x),
            PptxInches(y),
            PptxInches(width),
            PptxInches(height),
        )
        fill = style.get("fill")
        if _is_transparent(fill):
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _pptx_rgb(fill, "E2E8F0")
    stroke = style.get("stroke")
    if _is_transparent(stroke):
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _pptx_rgb(stroke, "1E293B")
        shape.line.width = PptxPt(max(0.25, float(style.get("strokeWidth") or 1) * 0.75))
    _set_shape_rotation(shape, block)


def _native_to_pptx(content: dict[str, Any], title: str) -> bytes:
    presentation = Presentation()
    format_name = content.get("format")
    canvas_width, canvas_height = (800.0, 600.0) if format_name == "4:3" else (960.0, 540.0)
    presentation.slide_width = PptxInches(10 if format_name == "4:3" else 13.333)
    presentation.slide_height = PptxInches(7.5)
    slide_width_inches = float(presentation.slide_width) / 914400
    slide_height_inches = float(presentation.slide_height) / 914400
    scale_x = slide_width_inches / canvas_width
    scale_y = slide_height_inches / canvas_height
    presentation.slides._sldIdLst.clear()
    native_slides = content.get("slides") if isinstance(content.get("slides"), list) else None
    if native_slides is None:
        if content.get("type") == "sheet_v1":
            lines = ["\t".join(str(value) for value in row) for row in (content.get("data") or [])]
        else:
            soup = BeautifulSoup(str(content.get("htmlContent") or content.get("content") or ""), "html.parser")
            lines = [node.get_text(" ", strip=True) for node in soup.find_all(["h1", "h2", "h3", "p", "li"]) if node.get_text(" ", strip=True)]
        chunks = [lines[index:index + 8] for index in range(0, len(lines), 8)] or [[title]]
        native_slides = [{"title": title if index == 0 else f"{title} · {index + 1}", "blocks": [{"type": "text", "content": "\n".join(chunk), "x": 70, "y": 100, "width": 820, "height": 360, "style": {"fontSize": 24}}]} for index, chunk in enumerate(chunks)]
    for index, native_slide in enumerate(native_slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = native_slide.get("backgroundColor")
        if not _is_transparent(background):
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _pptx_rgb(background, "FFFFFF")
        blocks = sorted(native_slide.get("blocks") or [], key=lambda block: float(block.get("zIndex", 0) or 0))
        for block in blocks:
            block_x, block_y, block_width, block_height = _pptx_box(block, canvas_width, canvas_height)
            x, y = block_x * scale_x, block_y * scale_y
            width, height = block_width * scale_x, block_height * scale_y
            if block.get("type") == "text":
                _add_text_to_slide(
                    slide,
                    str(block.get("content") or ""),
                    x, y, width, height,
                    block.get("style") or {},
                    scale_x, scale_y,
                    float(block.get("rotation") or 0),
                )
            elif block.get("type") == "image":
                match = re.match(r"data:([^;]+);base64,(.+)", str(block.get("content") or ""), re.DOTALL)
                if match:
                    try:
                        image = _cover_image_bytes(base64.b64decode(match.group(2)), block_width, block_height)
                        picture = slide.shapes.add_picture(io.BytesIO(image), PptxInches(x), PptxInches(y), PptxInches(width), PptxInches(height))
                        _set_shape_rotation(picture, block)
                    except Exception:
                        pass
            elif block.get("type") in {"rectangle", "ellipse", "line"}:
                _add_shape_to_slide(slide, block, x, y, width, height)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _native_to_xlsx(content: dict[str, Any], title: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "_", title)[:31] or "Foglio"
    rows = content.get("data") if isinstance(content.get("data"), list) else []
    for row in rows:
        sheet.append(list(row))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_document(content_json: str, title: str, target_format: str) -> ExportedDocument:
    target = target_format.lower().lstrip(".")
    if target not in SUPPORTED_EXPORT_FORMATS:
        raise ValueError(f"Formato di esportazione .{target} non supportato")
    try:
        content = json.loads(content_json)
    except (TypeError, ValueError) as exc:
        raise ValueError("Contenuto documento non valido") from exc
    safe_name = re.sub(r"[^\w\-. ]+", "_", title, flags=re.UNICODE).strip(" .") or "Documento"
    if target == "xlsx":
        data = _native_to_xlsx(content, title)
    elif target in {"ppt", "pptx"}:
        pptx = _native_to_pptx(content, title)
        data = pptx if target == "pptx" else _run_libreoffice(pptx, "pptx", "ppt")
    elif target in {"doc", "docx"}:
        docx = _native_to_docx(content, title)
        data = docx if target == "docx" else _run_libreoffice(docx, "docx", "doc")
    else:
        if content.get("type") == "presentation_v2" or isinstance(content.get("slides"), list):
            native = _native_to_pptx(content, title)
            data = _run_libreoffice(native, "pptx", "pdf")
        elif content.get("type") == "sheet_v1" or isinstance(content.get("data"), list):
            native = _native_to_xlsx(content, title)
            data = _run_libreoffice(native, "xlsx", "pdf")
        else:
            native = _native_to_docx(content, title)
            data = _run_libreoffice(native, "docx", "pdf")
    mime_type = MIME_BY_EXTENSION.get(target) or mimetypes.guess_type(f"file.{target}")[0] or "application/octet-stream"
    return ExportedDocument(f"{safe_name}.{target}", mime_type, data)
