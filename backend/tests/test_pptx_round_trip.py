"""PowerPoint ↔ slide editor round trips.

Decks are authored with python-pptx the way PowerPoint writes them (placeholders that inherit
their formatting from layout/master, theme colours, bullets with levels, groups, tables, notes),
imported into the editor model, exported back and re-imported.
"""

import base64
import io
import json

from bs4 import BeautifulSoup
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.services.document_conversion import export_document, import_document


def _deck() -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    # Slide 1: title + bullets from the "Title and Content" layout (formatting inherited from master).
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Fotosintesi"
    body = slide.placeholders[1].text_frame
    body.text = "Luce"
    second = body.add_paragraph()
    second.text = "Clorofilla"
    second.level = 1
    third = body.add_paragraph()
    third.text = "Glucosio"
    slide.notes_slide.notes_text_frame.text = "Ricordare l'equazione."

    # Slide 2: rich text box, shapes, group, table, background colour.
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xF4, 0xCC)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1.5))
    paragraph = box.text_frame.paragraphs[0]
    plain = paragraph.add_run()
    plain.text = "Testo normale più lungo "
    plain.font.size = Pt(24)
    bold = paragraph.add_run()
    bold.text = "grassetto rosso"
    bold.font.size = Pt(24)
    bold.font.bold = True
    bold.font.color.rgb = RGBColor(0xFE, 0x00, 0x4D)
    themed = paragraph.add_run()
    themed.text = " tema"
    themed.font.size = Pt(24)
    themed.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1

    rectangle = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(3), Inches(3), Inches(2))
    rectangle.fill.solid()
    rectangle.fill.fore_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    rectangle.line.color.rgb = RGBColor(0x10, 0x20, 0x30)
    rectangle.text_frame.text = "Dentro"
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5), Inches(6), Inches(8), Inches(6))

    group = slide.shapes.add_group_shape()
    circle = group.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9), Inches(1), Inches(1), Inches(1))
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(0x00, 0x80, 0x00)

    table = slide.shapes.add_table(2, 2, Inches(8), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _import(data: bytes) -> dict:
    return json.loads(import_document("deck.pptx", data).content_json)


def _text_blocks(slide: dict) -> list[dict]:
    return [block for block in slide["blocks"] if block["type"] == "text"]


def test_import_resolves_placeholders_bullets_and_notes():
    content = _import(_deck())
    assert content["format"] == "16:9"
    first = content["slides"][0]
    assert first["title"] == "Fotosintesi"
    assert first["notes"] == "Ricordare l'equazione."
    title_block = next(block for block in _text_blocks(first) if "Fotosintesi" in block["content"])
    # Master title style is 44pt → 44 canvas px on a 13.333in / 960px slide.
    assert abs(title_block["style"]["fontSize"] - 44) < 1
    body_block = next(block for block in _text_blocks(first) if "Luce" in block["content"])
    soup = BeautifulSoup(body_block["content"], "html.parser")
    outer = soup.find("ul")
    assert outer is not None, "body placeholder paragraphs inherit bullets from the master"
    nested = outer.find("ul")
    assert nested is not None and "Clorofilla" in nested.get_text()
    assert "Glucosio" in outer.get_text() and "Glucosio" not in nested.get_text()


def test_import_keeps_run_formatting_theme_colours_shapes_groups_and_tables():
    content = _import(_deck())
    slide = content["slides"][1]
    assert slide["backgroundColor"].lower() == "#fff4cc"
    rich = next(block for block in _text_blocks(slide) if "grassetto" in block["content"])
    soup = BeautifulSoup(rich["content"], "html.parser")
    strong = soup.find("strong")
    assert strong and "grassetto rosso" in strong.get_text()
    assert "#fe004d" in str(strong).lower()
    themed = soup.find(string=lambda s: s and "tema" in s).find_parent("span")
    assert themed is not None and "color" in themed.get("style", ""), "theme colour must resolve to a hex value"
    assert abs(rich["style"]["fontSize"] - 24) < 0.6

    types = [block["type"] for block in slide["blocks"]]
    assert "rectangle" in types and "ellipse" in types and "line" in types
    rectangle = next(block for block in slide["blocks"] if block["type"] == "rectangle")
    assert rectangle["style"]["fill"].lower() == "#abcdef" and rectangle["style"]["stroke"].lower() == "#102030"
    assert any("Dentro" in block["content"] for block in _text_blocks(slide))
    ellipse = next(block for block in slide["blocks"] if block["type"] == "ellipse")
    assert abs(ellipse["x"] - 9 * 72) < 2, "grouped shapes keep their slide position"
    table_text = next(block for block in _text_blocks(slide) if "A" in block["content"] and "B" in block["content"])
    assert "1" in table_text["content"] and "2" in table_text["content"]


def test_export_writes_runs_bullets_notes_and_reimports_equivalently():
    content = _import(_deck())
    exported = export_document(json.dumps(content), "Deck", "pptx").content
    presentation = Presentation(io.BytesIO(exported))
    assert presentation.slides[0].notes_slide.notes_text_frame.text == "Ricordare l'equazione."

    rich_shape = next(
        shape for shape in presentation.slides[1].shapes
        if getattr(shape, "has_text_frame", False) and "grassetto" in shape.text_frame.text
    )
    runs = [run for paragraph in rich_shape.text_frame.paragraphs for run in paragraph.runs]
    bold_run = next(run for run in runs if "grassetto" in run.text)
    assert bold_run.font.bold and str(bold_run.font.color.rgb) == "FE004D"
    assert "<" not in rich_shape.text_frame.text, "HTML must never leak into slide text"
    assert abs(bold_run.font.size.pt - 24) < 0.6

    bullets_shape = next(
        shape for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and "Luce" in shape.text_frame.text
    )
    levels = {p.text: p.level for p in bullets_shape.text_frame.paragraphs}
    assert levels["Clorofilla"] == 1 and levels["Luce"] == 0
    first_ppr = bullets_shape.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    assert first_ppr.find(qn("a:buChar")) is not None

    again = _import(exported)
    assert [s["title"] for s in again["slides"]] == [s["title"] for s in content["slides"]]
    for original, reimported in zip(content["slides"], again["slides"]):
        original_text = sorted(BeautifulSoup(b["content"], "html.parser").get_text() for b in _text_blocks(original))
        reimported_text = sorted(BeautifulSoup(b["content"], "html.parser").get_text() for b in _text_blocks(reimported))
        assert original_text == reimported_text
        assert original.get("notes") == reimported.get("notes")
    body = next(b for b in _text_blocks(again["slides"][0]) if "Luce" in b["content"])
    assert BeautifulSoup(body["content"], "html.parser").find("ul").find("ul") is not None


def test_legacy_plain_text_blocks_still_export():
    native = {
        "type": "presentation_v2", "format": "16:9",
        "slides": [{"id": "s", "title": "S", "blocks": [{
            "id": "t", "type": "text", "content": "Riga uno\nRiga due",
            "x": 10, "y": 10, "width": 400, "height": 100, "style": {"fontSize": 20},
        }]}],
    }
    presentation = Presentation(io.BytesIO(export_document(json.dumps(native), "S", "pptx").content))
    shape = next(s for s in presentation.slides[0].shapes if getattr(s, "has_text_frame", False))
    assert [p.text for p in shape.text_frame.paragraphs] == ["Riga uno", "Riga due"]


def test_background_image_becomes_locked_full_slide_block():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    image = io.BytesIO()
    Image.new("RGB", (32, 18), (10, 120, 200)).save(image, format="PNG")
    picture = slide.shapes.add_picture(io.BytesIO(image.getvalue()), 0, 0)
    blip = picture._element.find(".//" + qn("a:blip"))
    relationship_id = blip.get(qn("r:embed"))
    # Move the picture into the slide background (<p:bg><p:bgPr><a:blipFill>…).
    c_sld = slide._element.find(qn("p:cSld"))
    bg = etree.SubElement(c_sld, qn("p:bg"))
    c_sld.remove(bg)
    c_sld.insert(0, bg)
    bg_pr = etree.SubElement(bg, qn("p:bgPr"))
    blip_fill = etree.SubElement(bg_pr, qn("a:blipFill"))
    new_blip = etree.SubElement(blip_fill, qn("a:blip"))
    new_blip.set(qn("r:embed"), relationship_id)
    etree.SubElement(etree.SubElement(blip_fill, qn("a:stretch")), qn("a:fillRect"))
    etree.SubElement(bg_pr, qn("a:effectLst"))
    picture._element.getparent().remove(picture._element)
    buffer = io.BytesIO()
    presentation.save(buffer)

    content = _import(buffer.getvalue())
    first = content["slides"][0]["blocks"][0]
    assert first["type"] == "image" and first.get("locked") is True
    assert first["x"] == 0 and first["y"] == 0 and first["width"] == 800
    assert base64.b64decode(first["content"].split(",", 1)[1])[:4] == b"\x89PNG"
