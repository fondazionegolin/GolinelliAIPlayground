import base64
import io
import json
import zipfile

import fitz
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PptxInches

from app.services.document_conversion import _run_libreoffice, export_document, import_document


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Titolo DOCX", 1)
    document.add_paragraph("Testo modificabile")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(6), PptxInches(1))
    box.text = "Prima slide"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _rich_docx_bytes() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Intestazione modificabile"
    document.sections[0].footer.paragraphs[0].text = "Piè di pagina modificabile"
    document.add_heading("Titolo strutturato", 2)

    paragraph = document.add_paragraph()
    bold = paragraph.add_run("Grassetto")
    bold.bold = True
    paragraph.add_run(" ")
    italic = paragraph.add_run("Corsivo")
    italic.italic = True
    paragraph.add_run(" ")
    underline = paragraph.add_run("Sottolineato")
    underline.underline = True
    paragraph.add_run(" ")
    strike = paragraph.add_run("Barrato")
    strike.font.strike = True
    paragraph.add_run(" ")
    superscript = paragraph.add_run("Apice")
    superscript.font.superscript = True
    paragraph.add_run(" ")
    subscript = paragraph.add_run("Pedice")
    subscript.font.subscript = True
    paragraph.add_run(" ")
    styled = paragraph.add_run("Colorato")
    styled.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    styled.font.highlight_color = WD_COLOR_INDEX.YELLOW
    styled.font.name = "Arial"
    styled.font.size = Pt(18)
    paragraph.add_run(" ")
    inherited_style = document.styles.add_style("Evidenza importata", WD_STYLE_TYPE.CHARACTER)
    inherited_style.font.color.rgb = RGBColor(0x65, 0x43, 0x21)
    inherited_style.font.italic = True
    inherited = paragraph.add_run("Stile ereditato")
    inherited.style = inherited_style

    image_paragraph = document.add_paragraph()
    image_paragraph.add_run("Prima immagine")
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
        "/x8AAusB9Y9Z24QAAAAASUVORK5CYII="
    )
    image_paragraph.add_run().add_picture(io.BytesIO(png), width=Inches(0.25))
    image_paragraph.add_run("Dopo immagine")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Colonna A"
    table.cell(0, 1).text = "Colonna B"
    table.cell(1, 0).text = "Valore 1"
    table.cell(1, 1).text = "Valore 2"

    annotated = document.add_paragraph()
    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), "0")
    annotated._p.append(start)
    annotated.add_run("Testo commentato")
    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), "0")
    annotated._p.append(end)

    insertion = OxmlElement("w:ins")
    insertion.set(qn("w:author"), "Docente")
    insertion.set(qn("w:date"), "2026-07-24T10:00:00Z")
    inserted_run = OxmlElement("w:r")
    inserted_text = OxmlElement("w:t")
    inserted_text.text = "Testo aggiunto"
    inserted_run.append(inserted_text)
    insertion.append(inserted_run)
    annotated._p.append(insertion)

    deletion = OxmlElement("w:del")
    deletion.set(qn("w:author"), "Docente")
    deletion.set(qn("w:date"), "2026-07-24T10:01:00Z")
    deleted_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "Testo eliminato"
    deleted_run.append(deleted_text)
    deletion.append(deleted_run)
    annotated._p.append(deletion)

    output = io.BytesIO()
    document.save(output)
    source = io.BytesIO(output.getvalue())
    enriched = io.BytesIO()
    comments_xml = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:comment w:id="0" w:author="Revisore" w:date="2026-07-24T09:00:00Z">
    <w:p><w:r><w:t>Commento da conservare</w:t></w:r></w:p>
  </w:comment>
</w:comments>"""
    with zipfile.ZipFile(source, "r") as original, zipfile.ZipFile(enriched, "w") as target:
        for item in original.infolist():
            target.writestr(item, original.read(item.filename))
        target.writestr("word/comments.xml", comments_xml)
    return enriched.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Nome", "Valore"])
    sheet.append(["A", 42])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_bytes() -> bytes:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Prima pagina PDF")
    output = pdf.tobytes()
    pdf.close()
    return output


def test_imports_supported_editable_formats_and_preserves_source_metadata():
    samples = {
        "lezione.docx": (_docx_bytes(), "document"),
        "lezione.pptx": (_pptx_bytes(), "presentation"),
        "dati.xlsx": (_xlsx_bytes(), "sheet"),
        "dati.csv": ("Nome;Valore\nAlfa;13\nBeta;5\n".encode(), "sheet"),
        "dispensa.pdf": (_pdf_bytes(), "document"),
        "note.md": (b"# Titolo MD\n\nTesto **forte**", "document"),
    }
    for filename, (payload, expected_type) in samples.items():
        imported = import_document(filename, payload, file_id="00000000-0000-0000-0000-000000000001")
        content = json.loads(imported.content_json)
        assert imported.doc_type == expected_type
        assert content["source"]["preservedOriginal"] is True
        assert content["source"]["extension"] == filename.rsplit(".", 1)[1]
    pdf_content = json.loads(import_document("dispensa.pdf", samples["dispensa.pdf"][0]).content_json)
    assert pdf_content["previewImage"].startswith("data:image/jpeg;base64,")


def test_csv_import_detects_delimiter_and_creates_editable_cells():
    imported = import_document("dati.csv", b'Nome,Nota,Valore\nAlfa,"testo, con virgola",13\n')
    content = json.loads(imported.content_json)
    assert imported.doc_type == "sheet"
    assert content["type"] == "sheet_v1"
    assert content["data"] == [["Nome", "Nota", "Valore"], ["Alfa", "testo, con virgola", "13"]]


def test_native_exports_produce_valid_office_files():
    native = import_document("lezione.docx", _docx_bytes()).content_json
    docx = export_document(native, "Lezione", "docx")
    pptx = export_document(native, "Lezione", "pptx")
    assert Document(io.BytesIO(docx.content)).paragraphs
    assert Presentation(io.BytesIO(pptx.content)).slides


def test_presentation_export_preserves_visual_elements_and_bounds():
    pixel = (
        "data:image/png;base64,"
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
        "/x8AAusB9Y9Z24QAAAAASUVORK5CYII="
    )
    native = json.dumps({
        "type": "presentation_v2",
        "format": "4:3",
        "slides": [{
            "id": "slide-1",
            "title": "Titolo non renderizzato automaticamente",
            "backgroundColor": "#fff4cc",
            "blocks": [
                {
                    "id": "text", "type": "text", "content": "Testo che deve restare nella slide",
                    "x": -40, "y": 30, "width": 900, "height": 90, "rotation": 5,
                    "style": {"fontSize": 24, "fontFamily": "Arial", "color": "#123456", "backgroundColor": "#ffffff", "fontWeight": "bold", "textAlign": "center", "padding": 8},
                },
                {"id": "image", "type": "image", "content": pixel, "x": 40, "y": 150, "width": 180, "height": 120, "style": {}},
                {"id": "rect", "type": "rectangle", "content": "", "x": 260, "y": 150, "width": 180, "height": 120, "style": {"fill": "#abcdef", "stroke": "#102030", "strokeWidth": 3, "cornerRadius": 8}},
                {"id": "ellipse", "type": "ellipse", "content": "", "x": 480, "y": 150, "width": 120, "height": 120, "style": {"fill": "#fedcba", "stroke": "#203040"}},
                {"id": "line", "type": "line", "content": "", "x": 620, "y": 330, "width": 140, "height": 80, "style": {"stroke": "#ff0000", "strokeWidth": 4}},
            ],
        }],
    })

    exported = export_document(native, "Visuale", "pptx")
    presentation = Presentation(io.BytesIO(exported.content))
    assert presentation.slide_width == PptxInches(10)
    assert presentation.slide_height == PptxInches(7.5)
    slide = presentation.slides[0]
    assert len(slide.shapes) == 5
    assert slide.background.fill.fore_color.rgb == RGBColor(0xFF, 0xF4, 0xCC)
    for shape in slide.shapes:
        assert shape.left >= 0
        assert shape.top >= 0
        assert shape.left + shape.width <= presentation.slide_width
        assert shape.top + shape.height <= presentation.slide_height
    text_box = next(shape for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    assert text_box.text == "Testo che deve restare nella slide"
    assert text_box.text_frame.paragraphs[0].runs[0].font.size == Pt(18)


def test_presentation_pdf_uses_the_complete_pptx_renderer():
    native = json.dumps({
        "type": "presentation_v2",
        "format": "16:9",
        "slides": [{
            "id": "slide-1",
            "title": "Slide 1",
            "backgroundColor": "#223344",
            "blocks": [{"id": "shape", "type": "rectangle", "content": "", "x": 80, "y": 80, "width": 300, "height": 180, "style": {"fill": "#ff0000", "stroke": "#ffffff"}}],
        }],
    })
    exported = export_document(native, "Visuale", "pdf")
    assert exported.content.startswith(b"%PDF")


def test_docx_import_preserves_editable_structure_styles_images_and_annotations():
    content = json.loads(import_document("completo.docx", _rich_docx_bytes()).content_json)
    imported_html = content["htmlContent"]
    soup = BeautifulSoup(imported_html, "html.parser")

    assert soup.find("h2").get_text() == "Titolo strutturato"
    assert soup.find("strong", string=lambda value: value and "Grassetto" in value)
    assert soup.find("em", string=lambda value: value and "Corsivo" in value)
    assert soup.find("u", string=lambda value: value and "Sottolineato" in value)
    assert soup.find("s", string=lambda value: value and "Barrato" in value)
    assert soup.find("sup", string=lambda value: value and "Apice" in value)
    assert soup.find("sub", string=lambda value: value and "Pedice" in value)
    assert "color:#123456" in imported_html
    assert "background-color:#ffff00" in imported_html
    assert "font-size:18pt" in imported_html
    assert "font-family:&#x27;Arial&#x27;" in imported_html
    assert "color:#654321" in imported_html
    assert soup.find("em", string=lambda value: value and "Stile ereditato" in value)

    image_position = imported_html.index("<img ")
    assert imported_html.index("Prima immagine") < image_position < imported_html.index("Dopo immagine")
    assert 'src="data:image/png;base64,' in imported_html
    assert 'width="24"' in imported_html

    assert 'data-docx-region="header"' in imported_html
    assert "Intestazione modificabile" in imported_html
    assert 'data-docx-region="footer"' in imported_html
    assert "Piè di pagina modificabile" in imported_html
    assert "<table>" in imported_html
    assert "Colonna A" in imported_html
    assert "Valore 2" in imported_html

    assert 'data-docx-comment="true"' in imported_html
    assert 'data-comment-author="Revisore"' in imported_html
    assert 'data-comment-text="Commento da conservare"' in imported_html
    assert 'data-revision-kind="insert"' in imported_html
    assert "Testo aggiunto" in imported_html
    assert 'data-revision-kind="delete"' in imported_html
    assert "Testo eliminato" in imported_html


def test_pdf_export_uses_office_bridge():
    native = import_document("lezione.docx", _docx_bytes()).content_json
    exported = export_document(native, "Lezione", "pdf")
    assert exported.content.startswith(b"%PDF")


def test_legacy_office_formats_are_imported_as_editable_drafts():
    native_document = import_document("lezione.docx", _docx_bytes()).content_json
    legacy_doc = export_document(native_document, "Lezione", "doc").content
    legacy_ppt = export_document(native_document, "Lezione", "ppt").content
    legacy_xls = _run_libreoffice(_xlsx_bytes(), "xlsx", "xls")
    assert import_document("lezione.doc", legacy_doc).doc_type == "document"
    assert import_document("lezione.ppt", legacy_ppt).doc_type == "presentation"
    assert import_document("dati.xls", legacy_xls).doc_type == "sheet"
